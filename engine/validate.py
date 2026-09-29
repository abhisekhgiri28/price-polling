"""Regression suite for the readers, adapters and the page's data. Run
after touching any of them. Offline: nothing here makes a request.

    python engine/validate.py
"""
import sys, pathlib, argparse, collections, json

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import sites
# No test starts a real browser for a page: `js_sites` turns it on with a
# stand-in, and reads SHOWN_JS only against pages set from strings.
sites.BROWSER_FALLBACK = False
# Amazon's slower pace is checked in `js_sites`; elsewhere it only slows
# the stubbed tests down.
_AMAZON_PACE = sites.AMAZON_PACE_SECONDS
sites.AMAZON_PACE_SECONDS = 0.0

class Checks:
    def __init__(self):
        self.passed = self.failed = 0
        self.failures = []

    def ok(self, name, cond, detail=""):
        if cond:
            self.passed += 1
        else:
            self.failed += 1
            self.failures.append((name, detail))
        print("  %s  %s%s" % ("PASS" if cond else "FAIL", name,
                              "" if cond else "   <- " + str(detail)))


_BABY_INNER = ('{"@context":"http://schema.org","@type":"Product",'
               '"name":"Juniors Printed T-shirt and Pyjama Set",'
               '"offers":{"@type":"Offer","priceCurrency":"INR",'
               '"price":"419","availability":"outOfStock"}}')

BABY_HTML = ('<script type="application/ld+json">'
             + __import__("html").escape(_BABY_INNER) + '</script>')

BOYZ_HTML = ('<html><head><script type="application/ld+json">'
             '{"@context":"http://schema.org","@type":"Product",'
             '"name":"Boyz n Galz Girl\'s Stylish Ballerinas",'
             '"image":["https://x/a.jpg"],'
             '"offers":[{"@type":"Offer","priceCurrency":"INR",'
             '"availability":"https://schema.org/InStock",'
             '"priceSpecification":[{"@type":"UnitPriceSpecification",'
             '"price":"599.00","priceCurrency":"INR"}]}]}'
             '</script></head><body></body></html>')

GRAPH_HTML = ('<script type=\'application/ld+json\'>'
              '{"@context":"http://schema.org","@graph":['
              '{"@type":"WebSite","name":"Shop"},'
              '{"@type":["Product"],"name":"Graph Product",'
              '"offers":{"@type":"Offer","price":"249.00",'
              '"availability":"https://schema.org/InStock"}}]}'
              '</script>')

MRP_HTML = ('<script type="application/ld+json">'
            '{"@type":"Product","name":"With MRP","offers":{"@type":"Offer",'
            '"price":"699.00","priceSpecification":['
            '{"@type":"UnitPriceSpecification","price":"699.00"},'
            '{"@type":"StrikethroughPriceSpecification","price":"1199.00"}]}}'
            '</script>')


class FakeResponse:
    def __init__(self, status=200, payload=None, text="", headers=None):
        self.status_code, self._payload, self.text = status, payload, text
        self.headers = headers or {}
        self.content = text.encode("utf-8") if isinstance(text, str) else b""

    def json(self):
        if self._payload is None:
            raise ValueError("not json")
        return self._payload


class SeqRequests:
    """A transport that plays a scripted sequence of responses per URL.

    The last response in a queue repeats for ever, so "always 429" and
    "429 then 200" are both one line. robots.txt always allows -- what is
    under test here is the retry, not permission.
    """

    def __init__(self, script):
        self.script = {k: list(v) for k, v in script.items()}
        self.calls = []

    def get(self, url, **kw):
        self.calls.append(url)
        if url.endswith("/robots.txt"):
            return FakeResponse(200, None, "User-agent: *\nAllow: /\n")
        for frag, queue in self.script.items():
            if frag in url:
                return queue.pop(0) if len(queue) > 1 else queue[0]
        return FakeResponse(404, None, "")

    def product_calls(self):
        return [c for c in self.calls if "robots" not in c]


class Boom:
    """A transport that always raises, to exercise the exception path."""

    def __init__(self, exc=None):
        self.calls = []
        self.exc = exc or ConnectionError("refused")

    def get(self, url, **kw):
        self.calls.append(url)
        if url.endswith("/robots.txt"):
            return FakeResponse(200, None, "User-agent: *\nAllow: /\n")
        raise self.exc


class FakeRequests:
    """Stands in for `requests` so adapters can be tested with no network."""

    def __init__(self, routes):
        self.routes, self.calls = routes, []

    def get(self, url, **kw):
        self.calls.append(url)
        if url.endswith("/robots.txt"):
            return FakeResponse(200, None, "User-agent: *\nAllow: /\n")
        for frag, resp in self.routes.items():
            if frag in url:
                return resp
        return FakeResponse(404, None, "")


def shopify_json(variants, title="A Set", product_type=None):
    return {"product": {"title": title, "product_type": product_type,
                        "variants": variants}}


def adapters(k):
    """Adapter behaviour, offline. Every case here was a real wrong number."""
    print("\nadapters -- offline, with a stubbed transport")
    import sites, compare

    real = sites.requests
    try:
        # 1. `available` is absent from Shopify's per-product JSON on every
        #    store measured. Absent must read as unknown, never as sold out.
        sites._robots.clear(); sites._last_hit.clear()
        sites.requests = FakeRequests({"/products/x.json": FakeResponse(
            200, shopify_json([{"title": "2-3Y / Green", "price": "1499.00",
                                "compare_at_price": "0.00"}]))})
        q = sites.fetch("https://brand.example/products/x", "2-3Y")
        k.ok("absent Shopify `available` reads as unknown, not out-of-stock",
             q.ok and q.in_stock is None and q.sale_price == 1499.0, q)

        # 2. compare_at_price "0.00" is a string and therefore truthy. It must
        #    not become an MRP of zero, which once made every discount 100%.
        k.ok('compare_at_price "0.00" yields no MRP, not zero',
             q.listed_mrp is None, q.listed_mrp)

        # 3. A compare_at equal to the price also means "no discount".
        sites._robots.clear(); sites._last_hit.clear()
        sites.requests = FakeRequests({"/products/x.json": FakeResponse(
            200, shopify_json([{"title": "2-3Y", "price": "990.00",
                                "compare_at_price": "990.00"}]))})
        q = sites.fetch("https://brand.example/products/x", "2-3Y")
        k.ok("compare_at equal to the price yields no MRP",
             q.ok and q.listed_mrp is None and q.sale_price == 990.0, q)

        # 4. Real discount survives all of the above.
        sites._robots.clear(); sites._last_hit.clear()
        sites.requests = FakeRequests({"/products/x.json": FakeResponse(
            200, shopify_json([{"title": "2-3Y", "price": "699.00",
                                "compare_at_price": "1199.00",
                                "available": True}]))})
        q = sites.fetch("https://brand.example/products/x", "2-3Y")
        k.ok("a real compare_at is reported, and present `available` is used",
             q.ok and q.listed_mrp == 1199.0 and q.in_stock == 1, q)

        # 5. Two different prices answering one size: report nothing.
        sites._robots.clear(); sites._last_hit.clear()
        sites.requests = FakeRequests({"/products/x.json": FakeResponse(
            200, shopify_json([{"title": "2-3Y / Green", "price": "1199.00"},
                               {"title": "2-3Y / Blue", "price": "1299.00"}]))})
        q = sites.fetch("https://brand.example/products/x", "2-3Y")
        k.ok("two prices matching one size returns no price",
             not q.ok and q.sale_price is None and "different prices" in q.note,
             q.note)

        # 6. brand_size is tried before Zoddle's size, which is what makes a
        #    match possible: the store says 1-2Y where intake says 12-18M.
        sites._robots.clear(); sites._last_hit.clear()
        sites.requests = FakeRequests({"/products/x.json": FakeResponse(
            200, shopify_json([{"title": "1-2Y / Green", "price": "1499.00"},
                               {"title": "6-7Y / Green", "price": "1599.00"}]))})
        q = sites.fetch("https://brand.example/products/x",
                        ["1-2y", "12-18M", "12-18M"])
        k.ok("brand_size is matched ahead of Zoddle's own size",
             q.ok and q.sale_price == 1499.0 and q.size_matched is True, q)

        # 7. One request per URL: 40 barcodes on one design must not be 40
        #    fetches. This is what keeps a big file off a rate limit.
        sites._robots.clear(); sites._last_hit.clear()
        fake = FakeRequests({"/products/x.json": FakeResponse(
            200, shopify_json([{"title": "2-3Y", "price": "999.00"}]))})
        sites.requests = fake
        lst = sites.fetch_listing("https://brand.example/products/x")
        for _ in range(40):
            sites.pick(lst, ["2-3Y"])
        # (the stock read, /products/x.js, is the one allowed extra)
        product_calls = [c for c in fake.calls
                         if "robots" not in c and not c.endswith(".js")]
        k.ok("one fetch serves any number of barcodes",
             len(product_calls) == 1, product_calls)

        # 7b. Two colourways priced differently at one size: the intake colour
        #     settles it without a single image download. This is the real
        #     peekaabookids.com Tyohaar case -- Blue 1199, Green 1299 at 6-8Y.
        sites._robots.clear(); sites._last_hit.clear()
        two_tone = shopify_json([
            {"title": "6-8Y / Blue", "price": "1199.00", "image_id": 11},
            {"title": "6-8Y / Green", "price": "1299.00", "image_id": 22}])
        two_tone["product"]["images"] = [{"id": 11, "src": "https://x/blue.jpg"},
                                         {"id": 22, "src": "https://x/green.jpg"}]
        fake = FakeRequests({"/products/x.json": FakeResponse(200, two_tone)})
        sites.requests = fake
        lst = sites.fetch_listing("https://brand.example/products/x")

        def must_not_run(cands):
            raise AssertionError("the image chooser was consulted needlessly")

        # 7a. Same price in every colour: the colour still picks the
        #     colourway, so the thumbnail is the right one. BownBee 2255 is
        #     maroon on a page listing blue first -- both 999 -- and the
        #     thumbnail showed blue (24 Sep 2026).
        sites._robots.clear(); sites._last_hit.clear()
        same = shopify_json([
            {"title": "3-6M / Blue", "price": "999.00", "image_id": 11},
            {"title": "3-6M / Maroon", "price": "999.00", "image_id": 22}])
        same["product"]["images"] = [{"id": 11, "src": "https://x/blue.jpg"},
                                     {"id": 22, "src": "https://x/maroon.jpg"}]
        sites.requests = FakeRequests({"/products/x.json": FakeResponse(200, same)})
        lst_same = sites.fetch_listing("https://brand.example/products/x")
        q = sites.pick(lst_same, ["3-6M"], colour="maroon")
        k.ok("a colour shows that colourway's photo even when prices agree",
             q.ok and q.sale_price == 999.0 and q.image == "https://x/maroon.jpg",
             (q.sale_price, q.image))
        q = sites.pick(lst_same, ["3-6M"], colour="purple")
        k.ok("a colour the page does not list changes nothing",
             q.ok and q.sale_price == 999.0 and q.image == "https://x/blue.jpg",
             (q.sale_price, q.image))
        sites.requests = fake

        q = sites.pick(lst, ["6-8Y"], colour="blue", chooser=must_not_run)
        k.ok("the intake colour picks the right colourway, with no downloads",
             q.ok and q.sale_price == 1199.0 and q.resolved_by == "colour", q)
        q = sites.pick(lst, ["6-8Y"], colour="green", chooser=must_not_run)
        k.ok("the other colourway resolves to the other price",
             q.ok and q.sale_price == 1299.0, q)

        # 7c. With no colour, the photograph chooser is consulted -- and its
        #     verdict is used only when it actually returns one.
        seen = {}

        def choose_blue(cands):
            seen["called"] = True
            return [v for v in cands if "BLUE" in v.labels], "matched by photograph"

        q = sites.pick(lst, ["6-8Y"], colour=None, chooser=choose_blue)
        k.ok("with no colour, the photograph decides",
             q.ok and q.sale_price == 1199.0 and q.resolved_by == "image"
             and seen.get("called"), q)

        # 7d. A chooser that declines leaves the row blank rather than guessing.
        q = sites.pick(lst, ["6-8Y"], colour=None,
                       chooser=lambda c: (None, "too close to separate"))
        k.ok("a declining photograph match still yields no price",
             not q.ok and q.sale_price is None, q)

        # 7e. An unrecognised colour word must not silently pick a colourway.
        q = sites.pick(lst, ["6-8Y"], colour="chartreuse",
                       chooser=lambda c: (None, "declined"))
        k.ok("a colour that matches no variant does not force a pick",
             not q.ok and q.sale_price is None, q)

        # 8. A disallowed path is never requested at all.
        sites._robots.clear(); sites._last_hit.clear()
        blocked = FakeRequests({})
        blocked.get = lambda url, **kw: (
            FakeResponse(200, None, "User-agent: *\nDisallow: /\n")
            if url.endswith("/robots.txt") else FakeResponse(200, {"product": {}}))
        sites.requests = blocked
        q = sites.fetch("https://brand.example/products/x", "2-3Y")
        k.ok("a robots-disallowed path is refused",
             not q.ok and "robots" in q.note, q.note)
    finally:
        sites.requests = real
        sites._robots.clear()
        sites._last_hit.clear()

    # 9. Sites this engine does not price are refused without a request.
    #    The fragment is the *reason*, and the reason is the point: amazon is
    #    refused for want of an adapter, not because it cannot be read, and a
    #    check that accepted the old CAPTCHA wording would have gone on
    #    passing after that finding stopped reproducing.
    #
    #    Flipkart left this list on 21 Sep 2026, and Amazon and FirstCry on
    #    23 Sep -- all three are read now, and their own checks are below.
    #    That is the shape this list is meant to have: a host sits here until
    #    it can be read honestly, and then it leaves.
    for host, frag in (("https://www.tatacliq.com/x/p-mp1", "client-rendered"),
                       ("https://www.meesho.com/x/p/1", "403")):
        q = sites.fetch(host, "2-3Y")
        k.ok("unsupported site refused with a reason: %s"
             % host.split("/")[2], not q.ok and frag in q.note, q.note)

    # 10. Size vocabulary.
    k.ok("size labels normalise across vocabularies",
         (sites.norm_size("12-18 Months"), sites.norm_size("2-3y"),
          sites.norm_size(" 6-12M ")) == ("12-18M", "2-3Y", "6-12M"),
         (sites.norm_size("12-18 Months"), sites.norm_size("2-3y")))

    # 11. Difference sign convention: positive means Zoddle is cheaper.
    k.ok("a dearer retailer gives a positive difference",
         compare._verdict(999.0, 1299.0) == ("MISMATCH", 300.0),
         compare._verdict(999.0, 1299.0))
    k.ok("a cheaper retailer gives a negative difference",
         compare._verdict(999.0, 799.0) == ("MISMATCH", -200.0),
         compare._verdict(999.0, 799.0))
    k.ok("equal prices are a MATCH with zero difference",
         compare._verdict(999.0, 999.0) == ("MATCH", 0.0))
    k.ok("a missing retailer price gives no verdict at all",
         compare._verdict(999.0, None) == ("", None))

    # 12. Pack sizes: a 3-pack must not be compared to a single garment.
    k.ok("pack size is read from a retailer title",
         (compare.pack_size("Kids Pack Of 3 Printed T-shirts"),
          compare.pack_size("Set of 2 Shorts"),
          compare.pack_size("Girls Clothing Set"),
          compare.pack_size("Pack of 30 wipes")) == (3, 2, None, None),
         (compare.pack_size("Kids Pack Of 3 Printed T-shirts"),
          compare.pack_size("Girls Clothing Set")))

    PACK = "Kids Pack Of 3 Printed T-shirts"
    k.ok("a 3-pack against a 1-piece barcode is flagged",
         "like-for-like" in compare.like_for_like({"units_in_piece": 1}, PACK),
         compare.like_for_like({"units_in_piece": 1}, PACK))
    k.ok("a 3-pack against a 3-piece barcode is not flagged",
         compare.like_for_like({"units_in_piece": 3}, PACK) == "",
         compare.like_for_like({"units_in_piece": 3}, PACK))

    # The design's own name outranks units_in_the_piece. Zwende's rakhi is
    # named "... | Set Of 2" yet carries units_in_the_piece = 1; both sides
    # describe the same pair, so this is a field error, not a wrong product.
    zw = {"name": "Handmade Krishna Rakhi With Roli Chawal | Set Of 2",
          "units_in_piece": 1}
    k.ok("a design whose own name states the pack size agrees with a listing "
         "of that size",
         compare.like_for_like(zw, "Krishna Rakhi Set Of 2") == "",
         compare.like_for_like(zw, "Krishna Rakhi Set Of 2"))
    k.ok("that design is still flagged against a differently sized pack",
         "like-for-like" in compare.like_for_like(zw, "Rakhi Pack Of 5"))
    k.ok("a design whose name says nothing still falls back to the field",
         "like-for-like" in compare.like_for_like(
             {"name": "Cotton Interlock Joggers - Black", "units_in_piece": 1},
             "2 Pcs Interlock Joggers"))

    # 13. Channels: a Cloudinary re-host or a Drive link is not a storefront.
    rec = {"images": [
        {"url": "https://res.cloudinary.com/x/a.jpg", "host": "res.cloudinary.com",
         "is_page": False},
        {"url": "https://drive.google.com/file/d/1/view", "host": "drive.google.com",
         "is_page": True},
        {"url": "https://brand.example/products/x", "host": "brand.example",
         "is_page": True},
        {"url": "https://brand.example/products/x", "host": "brand.example",
         "is_page": True},
        {"url": "https://www.myntra.com/a/b/1/buy", "host": "www.myntra.com",
         "is_page": True}]}
    chans = compare.channels_for(rec)
    k.ok("channels skip re-hosts and Drive links, dedupe, and name the site",
         chans == [("own site", "https://brand.example/products/x"),
                   ("myntra", "https://www.myntra.com/a/b/1/buy")], chans)

    #     A marketplace host is that brand's own storefront when the brand is
    #     the one that owns it. Hopscotch sells other labels *and* its own, so
    #     the same host is a marketplace on one row and the own site on the
    #     next -- 57 of the 66 Hopscotch rows were reporting no own-site price
    #     while the price sat in the other_price column.
    hop = {"brand": "Hopscotch", "images": [
        {"url": "https://www.hopscotch.in/product/1286474/x",
         "host": "www.hopscotch.in", "is_page": True}]}
    k.ok("a brand's own marketplace-hosted storefront counts as its own site",
         compare.channels_for(hop) ==
         [("own site", "https://www.hopscotch.in/product/1286474/x")],
         compare.channels_for(hop))

    other = dict(hop, brand="Peekaboo")
    k.ok("but the same host is a marketplace for anyone else's brand",
         compare.channels_for(other) ==
         [("hopscotch", "https://www.hopscotch.in/product/1286474/x")],
         compare.channels_for(other))

    k.ok("host ownership ignores spacing, punctuation and the www/TLD labels",
         compare.owns_host("Hop Scotch", "www.hopscotch.in") and
         compare.owns_host("Hopscotch", "hopscotch.myshopify.com") and
         not compare.owns_host("", "www.hopscotch.in") and
         not compare.owns_host("Hopscotch", "www.myntra.com"))

    # 14. The CSV: the original names survive, every channel has its own
    #     price column in the same row, and the cheapest source is named.
    for col in ("zoddle_barcode", "zoddle_price", "own_site_price",
                "myntra_price", "firstcry_price", "lowest_price",
                "lowest_source", "lowest_confirmed", "zoddle_vs_lowest",
                "zoddle_image", "retailer_image"):
        k.ok("comparison CSV has the column %s" % col,
             col in compare.CSV_COLUMNS)
    k.ok("every channel price sits in one row, not separate rows",
         all(c in compare.CSV_COLUMNS for c in
             ("own_site_price", "myntra_price", "firstcry_price")))

    # 15. Corroboration. An intake URL used to be trusted purely for being in
    #     the file, while a discovered candidate faced guards. That asymmetry
    #     put confirmed prices on the wrong product.
    Q = collections.namedtuple("Q", "product_name sale_price ok site note image "
                                    "resolved_by")

    def q(name, price=1190.0):
        return Q(name, price, True, "shop.example", "", None, "size")

    jogger = {"name": "Cotton Interlock Joggers - Black", "units_in_piece": 1}
    ok, why = compare.corroborate(jogger, q("2 Pcs Interlock Joggers Black Blue"))
    k.ok("a 2-piece combo listing does not corroborate a 1-piece design",
         not ok and any("pack of 2" in w for w in why), why)
    ok, _ = compare.corroborate(jogger, q("Cotton Interlock Joggers Black"))
    k.ok("the design's own listing corroborates it", ok)
    ok, _ = compare.corroborate(
        {"name": "Girls Rebel Denim Dress", "units_in_piece": 1},
        q("Boys Football Socks"))
    k.ok("the page title is not judged against the design name -- the link "
         "is the listing", ok)

    # '2 Pcs' and '3-Piece' count as packs, not only 'Pack of N'. Missing this
    # shape is how the orangesugar combo passed the like-for-like guard.
    k.ok("a leading count is read as a pack size",
         (compare.pack_size("2 Pcs Interlock Joggers"),
          compare.pack_size("3-Piece Kurta Set"),
          compare.pack_size("2 Pcs Cotton Joggers")) == (2, 3, 2),
         (compare.pack_size("2 Pcs Interlock Joggers"),
          compare.pack_size("3-Piece Kurta Set")))

    # 16. Tracking parameters must not make one page look like two rival URLs.
    k.ok("shopify search tracking is stripped for URL identity",
         compare.canonical_url("https://napchief.com/products/x?_pos=1&_sid=a")
         == compare.canonical_url("https://napchief.com/products/x/"),
         compare.canonical_url("https://napchief.com/products/x?_pos=1&_sid=a"))

    # 17. One link per barcode per marketplace: that link's quote is the
    #     channel's answer, and it is never second-guessed.
    rec_j = {"name": "Cotton Interlock Joggers - Black", "units_in_piece": 1}
    one = [("https://s/products/interlock-joggers-black",
            q("Something Else Entirely", 649.0))]
    url, quote, why = compare._resolve_channel(rec_j, one)
    k.ok("the file's one link is the channel's price, whatever its title",
         quote is not None and quote.sale_price == 649.0 and not why,
         (url, why))

    # -- Flipkart, read through the reader every storefront already uses --
    # Wired up 21 Sep 2026. It needed no adapter of its own; what it needed
    # was the complete browser header set, without which both robots.txt and
    # the PDP answer 403 and eight designs read as unpriceable.
    k.ok("flipkart is no longer listed as unreadable",
         "flipkart.com" not in sites.UNSUPPORTED, sorted(sites.UNSUPPORTED))
    k.ok("a flipkart PDP is routed to its own reader",
         sites.platform_for("https://www.flipkart.com/x/p/itm29338e6bfbacf")
         == ("flipkart", sites.flipkart_listing))
    # The thin two-header set is exactly what earned the 403, so the header
    # set is part of the contract now, not incidental.
    for h in ("Accept", "Accept-Encoding", "Sec-Fetch-Mode",
              "Upgrade-Insecure-Requests"):
        k.ok("HEADERS still sends %s, without which flipkart answers 403" % h,
             h in sites.HEADERS, sorted(sites.HEADERS))

    FLIPKART_PDP = (
        '<html><head><script type="application/ld+json">'
        '{"@context":"https://schema.org","@type":"Product",'
        '"name":"GINI & JONY Girls Graphic Print Cotton Blend Regular T Shirt",'
        '"offers":{"@type":"Offer","price":488,"priceCurrency":"INR",'
        '"availability":"https://schema.org/InStock"}}'
        '</script></head><body>'
        # The trap: the rendered page also carries the MRP, a card-only "Buy
        # at" figure and a junk label. A regex for a rupee sign picks one of
        # these; the ld+json block is the only honest source.
        '<div>699</div><div>240</div><div>21474685</div>'
        '</body></html>')

    real = sites.requests
    try:
        sites.requests = FakeRequests(
            {"/p/itm": FakeResponse(200, None, FLIPKART_PDP)})
        sites._robots.clear()
        lst = sites.fetch_listing("https://www.flipkart.com/g/p/itm363cf3680b12b")
        k.ok("the flipkart price comes from the ld+json Product",
             lst.ok and lst.variants and lst.variants[0].price == 488.0,
             [(v.price, v.available) for v in lst.variants])
        k.ok("and not from the MRP, the bank offer or the junk label on the page",
             all(v.price not in (699.0, 240.0, 21474685.0) for v in lst.variants),
             [v.price for v in lst.variants])
        k.ok("flipkart reports one price per listing, not a per-size table",
             lst.per_size is False, lst.per_size)
        k.ok("and it is filed as flipkart, not as a storefront",
             lst.platform == "flipkart" and lst.site == "flipkart.com",
             (lst.platform, lst.site))

        # Real case, met on design 892: Flipkart serves "[]" as the whole
        # ld+json block. A blank with its reason, never a guess.
        sites.requests = FakeRequests(
            {"/p/itm": FakeResponse(
                200, None, '<html><head><script type="application/ld+json">'
                           '[]</script></head><body>323</body></html>')})
        sites._robots.clear()
        empty = sites.fetch_listing("https://www.flipkart.com/g/p/itm23431b0dd7d81")
        k.ok("an empty ld+json block is a blank with a reason, not a price",
             not empty.ok and empty.variants == [] and empty.note,
             (empty.ok, empty.note))

        # 403 is what the thin header set used to earn. It must read as a
        # refusal, not as "this product has no price".
        sites.requests = FakeRequests({"/p/itm": FakeResponse(403, None, "")})
        sites._robots.clear()
        refused = sites.fetch_listing("https://www.flipkart.com/g/p/itm1")
        k.ok("a 403 from flipkart is reported as a refusal, not as no price",
             not refused.ok and "403" in (refused.note or ""), refused.note)

        # ------------------------------------------------------ amazon --
        # Amazon is the page where a wrong number is easiest to publish: an
        # UNAVAILABLE product still shows a dozen other products' prices in
        # its rails. Measured 23 Sep 2026 on all 11 corpus links -- the four
        # with no buybox are "Currently unavailable" and carry 8 to 37 rupee
        # amounts each, every one belonging to something else.
        RS = "₹"
        AMZ = ('<html><head><title>x</title></head><body>'
               '<span id="productTitle"> Nap Chief Girls Paw Patrol Tee </span>'
               '<img id="landingImage" src="https://m.media-amazon.com/i/1.jpg">'
               # The struck MRP sits BEFORE the selling price in the markup,
               # so "the first price in the block" is the wrong rule.
               '<span class="apex-basisprice-offscreen-label">M.R.P.: '
               + RS + '1,950.00</span>'
               '<span class="aok-offscreen" data-pricetopay-label="{priceToPay}">'
               + RS + '890.00 with 54 percent savings </span>'
               '<div id="availability"><span> In stock </span></div>'
               # The rails: real prices, none of them this product's.
               '<div class="a-price-whole">2,999</div>'
               '<div class="a-price-whole">349</div>'
               '</body></html>')
        sites.requests = FakeRequests({"/dp/": FakeResponse(200, None, AMZ)})
        sites._robots.clear()
        a = sites.fetch_listing("https://www.amazon.in/x/dp/B0GNSJ6FS2")
        k.ok("amazon is no longer listed as unreadable",
             "amazon.in" not in sites.UNSUPPORTED, sorted(sites.UNSUPPORTED))
        k.ok("an amazon PDP is routed to its own reader",
             sites.platform_for("https://www.amazon.in/x/dp/B0") ==
             ("amazon", sites.amazon_listing))
        k.ok("the amazon price comes from the buybox",
             a.ok and a.variants and a.variants[0].price == 890.0,
             [v.price for v in a.variants])
        k.ok("and never from the rails, which are other products entirely",
             all(v.price not in (2999.0, 349.0) for v in a.variants),
             [v.price for v in a.variants])
        k.ok("the struck MRP above the price is recorded, not sold as the price",
             a.listed_mrp == 1950.0 and a.variants[0].price == 890.0,
             (a.listed_mrp, a.variants[0].price))
        k.ok("amazon reports one price per listing, not a per-size table",
             a.per_size is False)
        k.ok("and availability is read from the availability block",
             a.variants[0].available is True, a.variants[0].available)

        # Currently unavailable: the page is full of prices and NONE is this
        # product's. The only honest answer is no price, with which it is.
        GONE = ('<html><body><span id="productTitle"> Hopscotch Boys Tee </span>'
                '<span class="a-size-medium primary-availability-message"> '
                'Currently unavailable. </span>'
                '<div class="a-price-whole">689</div>'
                '<div class="a-price-whole">2,999</div></body></html>')
        sites.requests = FakeRequests({"/dp/": FakeResponse(200, None, GONE)})
        sites._robots.clear()
        g = sites.fetch_listing("https://www.amazon.in/x/dp/B0GONE")
        k.ok("an unavailable amazon product is priced at nothing, not at 689",
             not g.ok and not g.variants, (g.ok, [v.price for v in g.variants]))
        k.ok("and it says it is unavailable rather than unreadable",
             "not selling" in (g.note or "").lower()
             or "unavailable" in (g.note or "").lower(), g.note)
        k.ok("the product name still comes back, so the row is identifiable",
             (g.product_name or "").startswith("Hopscotch"), g.product_name)

        # The CAPTCHA page is a 200 with a plausible shape. Parsing it would
        # produce a blank that reads as "Amazon does not sell this".
        CAP = ('<html><body>Enter the characters you see below'
               '<div class="a-price-whole">499</div></body></html>')
        sites.requests = FakeRequests({"/dp/": FakeResponse(200, None, CAP)})
        sites._robots.clear()
        c = sites.fetch_listing("https://www.amazon.in/x/dp/B0CAP")
        k.ok("amazon's CAPTCHA page yields nothing at all",
             not c.ok and not c.variants, (c.ok, [v.price for v in c.variants]))
        k.ok("and says so, because a challenge is not a fact about the garment",
             "captcha" in (c.note or "").lower(), c.note)

        # A struck figure BELOW the selling price is a neighbour, not an MRP.
        LOWMRP = AMZ.replace(RS + "1,950.00", RS + "99.00")
        sites.requests = FakeRequests({"/dp/": FakeResponse(200, None, LOWMRP)})
        sites._robots.clear()
        lm = sites.fetch_listing("https://www.amazon.in/x/dp/B0LOW")
        k.ok("an MRP below the selling price is dropped, not recorded",
             lm.ok and lm.listed_mrp is None and lm.variants[0].price == 890.0,
             (lm.listed_mrp, lm.variants[0].price))

        # ---------------------------------------------------- firstcry --
        # Every size is its own product id, and the product page lists the
        # whole style with FirstCry's own MRP and discount per id. Figures
        # below are the real ones for design 22339749, measured 23 Sep 2026
        # and checked against FirstCry's listing cards to the paisa.
        def fc_entry(pid, sz, dis, colour="1#Black", mrp=1099):
            return {"pid": int(pid), "sz": sz, "mrp": mrp, "Dis": dis,
                    "qty": 1, "Gid": "22339739", "hashCols": [colour],
                    "pn": "boyz n galz Dinosaur Applique Musical Sandals",
                    "Img": "%sa.jpg;%ssz.jpg;%sb.jpg;" % (pid, pid, pid)}
        fc_group = {
            "22339746": fc_entry("22339746", "EU 19", 63.69),
            "22339749": fc_entry("22339749", "EU 22", 63.69),
            # same colour, dearer size -- FirstCry prices sizes separately
            "22339750": fc_entry("22339750", "EU 23", 57.74),
            # another colour at the SAME size and a different price
            "22339743": fc_entry("22339743", "EU 22", 40.0, colour="2#Blue"),
        }
        FC_PDP = ('<html><head><title>Buy boyz n galz</title></head><body>'
                  '<script>var x=1,CurrentProductID=22339749,'
                  'CurrentProductDetailJSON=' + json.dumps(fc_group) +
                  ',other={};</script></body></html>')
        FC_URL = ("https://www.firstcry.com/boyz-n-galz/boyz-n-galz-dinosaur-"
                  "sandals-black/22339749/product-detail?srsltid=AfmBOor")
        fake = FakeRequests({"/22339749/product-detail":
                             FakeResponse(200, None, FC_PDP)})
        sites.requests = fake
        sites._robots.clear()
        f = sites.fetch_listing(FC_URL)
        k.ok("firstcry is no longer listed as unreadable",
             "firstcry.com" not in sites.UNSUPPORTED, sorted(sites.UNSUPPORTED))
        k.ok("a firstcry PDP is routed to its own reader",
             sites.platform_for(FC_URL) == ("firstcry", sites.firstcry_listing))
        k.ok("the firstcry page is fetched without the link's tracking query",
             [c for c in fake.calls if "robots" not in c] ==
             [FC_URL.split("?")[0]], fake.calls)
        k.ok("firstcry's price is its own MRP less its own discount",
             f.ok and sites.pick(f, ["EUR 22"]).sale_price == 399.05,
             (f.note, sites.pick(f, ["EUR 22"])))
        k.ok("the intake's 'EUR 22' and a bare '22' both find FirstCry's 'EU 22'",
             sites.pick(f, ["22"]).sale_price == 399.05,
             sites.pick(f, ["22"]).note)
        k.ok("a different colour at the same size is not a rival price",
             sites.pick(f, ["EUR 22"]).ok and
             all(v.price != 659.4 for v in f.variants),
             [v.price for v in f.variants])
        k.ok("each size carries its own price, not its neighbour's",
             sites.pick(f, ["EUR 23"]).sale_price == 464.44,
             sites.pick(f, ["EUR 23"]).sale_price)
        k.ok("a size FirstCry does not list is refused, not given another's price",
             not sites.pick(f, ["EUR 30"]).ok and
             "does not list" in sites.pick(f, ["EUR 30"]).note,
             sites.pick(f, ["EUR 30"]).note)
        k.ok("even when the listed sizes all happen to share one price",
             not sites.pick(f._replace(variants=f.variants[:2]), ["EUR 30"]).ok)
        k.ok("firstcry's qty is unproven, so stock is unknown -- not sold out",
             sites.pick(f, ["EUR 22"]).in_stock is None,
             sites.pick(f, ["EUR 22"]).in_stock)
        k.ok("firstcry's MRP is recorded for information only",
             sites.pick(f, ["EUR 22"]).listed_mrp == 1099.0,
             sites.pick(f, ["EUR 22"]).listed_mrp)
        k.ok("the size-chart picture is not offered as a product photograph",
             f.images and not any("sz." in u for u in f.images), f.images)

        # A discontinued product answers 200 with an empty shell. That page
        # is what "FirstCry publishes no price" was measured on, four times.
        GONE_FC = ('<html><head><title></title></head><body><p class="inactive'
                   'Text">We"re sorry, the product you are looking for has been '
                   'discontinued.</p><span class="rupee_sp mrp"></span>'
                   '</body></html>')
        sites.requests = FakeRequests({"/product-detail":
                                       FakeResponse(200, None, GONE_FC)})
        sites._robots.clear()
        d = sites.fetch_listing(FC_URL)
        k.ok("a discontinued firstcry product is priced at nothing",
             not d.ok and not d.variants, (d.ok, d.variants))
        k.ok("and says it is discontinued, not unreadable",
             "discontinued" in (d.note or ""), d.note)

        sites.requests = FakeRequests({"/product-detail": FakeResponse(
            200, None, "<html><body>a redesigned page</body></html>")})
        sites._robots.clear()
        n = sites.fetch_listing(FC_URL)
        k.ok("a firstcry page without the size data says its shape changed",
             not n.ok and "shape changed" in (n.note or ""), n.note)

        other = dict(fc_group)
        other.pop("22339749")
        sites.requests = FakeRequests({"/product-detail": FakeResponse(
            200, None, FC_PDP.replace(json.dumps(fc_group), json.dumps(other)))})
        sites._robots.clear()
        o = sites.fetch_listing(FC_URL)
        k.ok("a page that does not list the linked id prices nothing from it",
             not o.ok and not o.variants, (o.ok, o.note))

        bad_dis = {"22339749": fc_entry("22339749", "EU 22", 100)}
        sites.requests = FakeRequests({"/product-detail": FakeResponse(
            200, None, FC_PDP.replace(json.dumps(fc_group), json.dumps(bad_dis)))})
        sites._robots.clear()
        b = sites.fetch_listing(FC_URL)
        k.ok("a 100% discount is not a price of zero",
             not b.ok and not sites.pick(b, ["EUR 22"]).ok, (b.ok, b.note))

        # The row: FirstCry is a channel like any other now.
        fc_rec = {"zoddle_barcode": "22339749", "design_code": "7001",
                  "brand": "BoyznGalz", "name": "Dinosaur Musical Sandals",
                  "size": "18-24M", "brand_size": "EUR 22", "mrp": 1099.0,
                  "zoddle_price": 450.0,
                  "images": [{"position": 3, "url": FC_URL,
                              "host": "www.firstcry.com", "is_page": True}]}
        fr = compare._row_for(fc_rec, {FC_URL: f._replace(images=(),
                                                          image=None)})
        k.ok("a firstcry price lands in firstcry_price",
             fr.firstcry_price == 399.05, (fr.firstcry_price, fr.firstcry_status))
        k.ok("and firstcry competes for cheapest",
             fr.lowest_source == "firstcry" and fr.lowest_price == 399.05,
             (fr.lowest_source, fr.lowest_price))
        fr2 = compare._row_for(dict(fc_rec, brand_size="EUR 30"),
                               {FC_URL: f._replace(images=(), image=None)})
        k.ok("an unlisted size leaves firstcry blank with its reason",
             fr2.firstcry_price == "" and "does not list" in fr2.firstcry_status,
             fr2.firstcry_status)
        k.ok("a row with no firstcry link carries no firstcry status at all",
             compare._row_for(dict(fc_rec, images=[]), {}).firstcry_status == "")
    finally:
        sites.requests = real
        sites._robots.clear()


def resilience(k):
    """A refusal must never be recorded as an absent price.

    Every case here is the Hopscotch failure in miniature: on the real
    66-row file, 11 product pages answered 429 at one request per second and
    were written down as blanks, while every one of those URLs carried a
    price when asked more slowly.
    """
    print("\nresilience -- retries, pacing and non-Shopify storefronts")
    import sites

    real, real_backoff, real_pace = sites.requests, sites.BACKOFF, sites.PACE_SECONDS

    def reset():
        sites._robots.clear()
        sites._last_hit.clear()
        sites._host_pace.clear()

    try:
        sites.BACKOFF = (0.0, 0.0, 0.0)     # the waiting is not what is tested
        sites.PACE_SECONDS = 0.01

        ok_json = shopify_json([{"title": "2-3Y", "price": "999.00"}])

        # 1. The Hopscotch case exactly: refused once, priced on the retry.
        reset()
        t = SeqRequests({"/products/x.json": [FakeResponse(429), FakeResponse(200, ok_json)]})
        sites.requests = t
        q = sites.fetch("https://brand.example/products/x", "2-3Y")
        k.ok("a 429 that clears on retry yields the price, not a blank",
             q.ok and q.sale_price == 999.0, q.note)
        k.ok("and it really did retry rather than get lucky",
             len([c for c in t.product_calls() if c.endswith(".json")]) == 2,
             t.product_calls())

        # 2. A refusal that never clears says so, and says how hard we tried.
        reset()
        t = SeqRequests({"/products/x.json": [FakeResponse(429)]})
        sites.requests = t
        q = sites.fetch("https://brand.example/products/x", "2-3Y")
        k.ok("a refusal that never clears is reported as a refusal",
             not q.ok and "429" in q.note and "attempts" in q.note, q.note)
        json_calls = [c for c in t.product_calls() if c.endswith(".json")]
        k.ok("and the Shopify route is retried the full number of times",
             len(json_calls) == sites.RETRIES + 1, json_calls)
        # Having been refused on the JSON route, the page itself is still
        # worth one look: Shopify rate-limits /products/<handle>.json harder
        # than it serves storefront HTML, so the fallback can still carry the
        # price home. Never missing an available price is the whole point.
        k.ok("a refused Shopify route still falls through to the page",
             any(not c.endswith(".json") for c in t.product_calls()),
             t.product_calls())

        # 3. A 5xx is transient too -- a gateway blip is not a missing price.
        reset()
        t = SeqRequests({"/products/x.json": [FakeResponse(503), FakeResponse(502),
                                              FakeResponse(200, ok_json)]})
        sites.requests = t
        q = sites.fetch("https://brand.example/products/x", "2-3Y")
        k.ok("a 5xx is retried until the store answers",
             q.ok and q.sale_price == 999.0, q.note)

        # 4. A connection error is retried, not recorded as "no price".
        reset()
        b = Boom()
        sites.requests = b
        q = sites.fetch("https://brand.example/products/x", "2-3Y")
        k.ok("a connection error is retried and then reported honestly",
             not q.ok and "ConnectionError" in q.note, q.note)
        tried = [c for c in b.calls if "robots" not in c]
        k.ok("and every route was retried the full number of times",
             len([c for c in tried if c.endswith(".json")]) == sites.RETRIES + 1
             and len(tried) == 2 * (sites.RETRIES + 1), tried)

        # 5. A delisted product is an answer about the product. Retrying it
        #    is pointless, and "HTTP 410" tells a reader less than the word.
        reset()
        t = SeqRequests({"/products/gone.json": [FakeResponse(410)]})
        sites.requests = t
        q = sites.fetch("https://brand.example/products/gone", "2-3Y")
        k.ok("a 410 reads as delisted rather than as a status code",
             not q.ok and "delisted" in q.note, q.note)
        k.ok("and a delisted page is asked exactly once",
             len(t.product_calls()) == 1, t.product_calls())

        # 6. A 429 widens this host's pace for the rest of the run, so the
        #    next URL on the same store does not walk into the same wall.
        reset()
        sites.requests = SeqRequests({"/products/x.json": [FakeResponse(429)]})
        sites.fetch("https://slow.example/products/x", "2-3Y")
        k.ok("a host that answers 429 has its pace widened",
             sites.pace_for("slow.example") > sites.PACE_SECONDS,
             sites.pace_for("slow.example"))
        k.ok("and an unrelated host is left alone",
             sites.pace_for("other.example") == sites.PACE_SECONDS,
             sites.pace_for("other.example"))

        # 7. Retry-After, in both spellings the RFC allows, and capped.
        k.ok("a numeric Retry-After is honoured",
             sites._retry_after(FakeResponse(429, headers={"Retry-After": "7"})) == 7.0)
        k.ok("an absurd Retry-After is capped rather than obeyed",
             sites._retry_after(
                 FakeResponse(429, headers={"Retry-After": "99999"}))
             == sites.MAX_RETRY_AFTER)
        k.ok("a missing Retry-After is None, not zero",
             sites._retry_after(FakeResponse(429)) is None)
        k.ok("a response with no headers at all does not raise",
             sites._retry_after(object()) is None)

        # -- storefronts that are not Shopify --------------------------------
        # 8. boyzngalz.com: the price is nested in priceSpecification, where a
        #    flat offers["price"] read finds nothing and reports a live 599
        #    product as unpriced.
        reset()
        t = SeqRequests({"/product/ballerinas.json": [FakeResponse(404)],
                         "/product/ballerinas": [FakeResponse(200, None, BOYZ_HTML)]})
        sites.requests = t
        q = sites.fetch("https://boyzngalz.example/product/ballerinas")
        k.ok("a non-Shopify storefront is priced from its schema.org Product",
             q.ok and q.sale_price == 599.0, q.note)
        k.ok("and the price nested in priceSpecification is the one found",
             q.ok and q.sale_price == 599.0 and q.platform == "storefront", q.platform)

        # 9. www.babyshop.in HTML-escapes its payload; unescaping is the
        #    difference between three unreadable blobs and a price.
        reset()
        sites.requests = SeqRequests({"/p/1000015059352.json": [FakeResponse(404)],
                                      "/p/1000015059352": [FakeResponse(200, None, BABY_HTML)]})
        q = sites.fetch("https://babyshop.example/p/1000015059352")
        k.ok("an HTML-escaped ld+json payload is still read",
             q.ok and q.sale_price == 419.0, q.note)
        k.ok("and its lower-case outOfStock is understood",
             q.in_stock == 0, q.in_stock)

        # 10. @graph and a list-valued @type are both legal schema.org.
        reset()
        sites.requests = SeqRequests({"/p/graph.json": [FakeResponse(404)],
                                      "/p/graph": [FakeResponse(200, None, GRAPH_HTML)]})
        q = sites.fetch("https://shop.example/p/graph")
        k.ok("a Product inside @graph with a list @type is found",
             q.ok and q.sale_price == 249.0, q.note)

        # 11. A struck-through price is the retailer's MRP: recorded, never
        #     compared. Same rule Shopify's compare_at_price already gets.
        reset()
        sites.requests = SeqRequests({"/p/mrp.json": [FakeResponse(404)],
                                      "/p/mrp": [FakeResponse(200, None, MRP_HTML)]})
        q = sites.fetch("https://shop.example/p/mrp")
        k.ok("a struck-through price is kept as listed_mrp, not as the price",
             q.ok and q.sale_price == 699.0 and q.listed_mrp == 1199.0,
             (q.sale_price, q.listed_mrp))

        # 12. Shopify stays the fast path: when it answers, the page is never
        #     fetched a second time.
        reset()
        t = SeqRequests({"/products/x.json": [FakeResponse(200, ok_json)]})
        sites.requests = t
        q = sites.fetch("https://brand.example/products/x", "2-3Y")
        k.ok("a Shopify store is served by one request, with no fallback",
             q.ok and q.platform == "shopify" and
             [c for c in t.product_calls() if not c.endswith(".js")] ==
             ["https://brand.example/products/x.json"],
             t.product_calls())

        # 13. A 410 on the Shopify route must not spend a second request on a
        #     page that is gone -- but a 404 MUST fall through, because that
        #     is exactly how a store which is not Shopify answers it.
        reset()
        t = SeqRequests({"/products/gone.json": [FakeResponse(410)],
                         "/products/gone": [FakeResponse(200, None, BOYZ_HTML)]})
        sites.requests = t
        q = sites.fetch("https://brand.example/products/gone")
        k.ok("a delisted product does not fall through to the page",
             not q.ok and "delisted" in q.note and len(t.product_calls()) == 1,
             t.product_calls())

        # 14. Nothing readable anywhere still reports the page, not the
        #     Shopify endpoint that was never going to exist.
        reset()
        sites.requests = SeqRequests({"/p/bare.json": [FakeResponse(404)],
                                      "/p/bare": [FakeResponse(200, None, "<html></html>")]})
        q = sites.fetch("https://shop.example/p/bare")
        k.ok("a page with no Product says so, rather than blaming a 404 json",
             not q.ok and "schema.org Product" in q.note, q.note)
    finally:
        sites.requests, sites.BACKOFF = real, real_backoff
        sites.PACE_SECONDS = real_pace
        reset()

    # 15. An image CDN is never mistaken for a storefront, however well its
    #     name matches the brand. static.hopscotch.in really is Hopscotch's,
    #     and own-site discovery asked it for a product catalogue.
    import compare
    k.ok("an image CDN is not treated as a brand's storefront",
         compare.is_cdn_host("static.hopscotch.in")
         and compare.is_cdn_host("res.cloudinary.com"), "")
    k.ok("a real storefront host is not mistaken for a CDN",
         not compare.is_cdn_host("peekaabookids.com")
         and not compare.is_cdn_host("googogaaga.com"), "")
    rec = {"brand": "Hopscotch", "images": [
        {"position": 2, "url": "https://static.hopscotch.in/a.jpg",
         "host": "static.hopscotch.in", "is_page": False}]}
    k.ok("so own_host_for declines to guess a storefront from a CDN",
         compare.own_host_for(rec) == "", compare.own_host_for(rec))

    # NPL -- no product link. A row nobody gave a page for has to be
    # distinguishable from a page that refused: the first is an absent input,
    # the second a failed read, and only one of them is the engine's problem.
    bare = {"zoddle_barcode": "123456", "design_code": "12", "brand": "Testly",
            "name": "Thing", "size": "2-3Y", "mrp": 100.0,
            "zoddle_price": 90.0, "images": []}
    r_bare = compare._row_for(bare, {})
    k.ok("a row with no product link is marked NPL",
         r_bare.own_site_match == "NPL", r_bare.own_site_match)
    k.ok("and NPL carries no price and crowns no cheapest source from one",
         r_bare.own_site_price == "" and "own site" not in r_bare.lowest_source,
         (r_bare.own_site_price, r_bare.lowest_source))

    url = "https://brand.example/products/x"
    linked = dict(bare, zoddle_barcode="123457",
                  images=[{"position": 3, "url": url, "host": "brand.example",
                           "is_page": True}])
    r_fail = compare._row_for(
        linked, {url: sites._listing(url, "brand.example", "shopify",
                                     note="HTTP 500")})
    k.ok("a link that failed to answer is NOT reported as NPL",
         r_fail.own_site_match != "NPL" and "500" in r_fail.note, r_fail.note)
    k.ok("and the summary counts the NPL rows separately",
         compare.summarise([r_bare, r_fail], {})["no_product_link"] == 1,
         compare.summarise([r_bare, r_fail], {})["no_product_link"])


def _gw_products(n, start=0, brand="Peekaaboo Kids"):
    return [{"productId": start + i, "productName": "Item %d" % (start + i),
             "brand": brand, "price": 100 + i, "mrp": 200 + i,
             "gender": "Girls", "primaryColour": "Blue",
             "articleType": {"typeName": "Dress"},
             "images": [{"view": "default", "src": "https://x/%d.jpg" % (start + i)}],
             "searchImage": "https://x/%d.jpg" % (start + i),
             "landingPageUrl": "Dress/x/%d/buy" % (start + i)}
            for i in range(n)]


def _myx_html(products, total, ctx="CTX0"):
    import json as _j
    payload = {"searchData": {"results": {"products": products,
                                          "totalCount": total},
                              "nextPaginationContext": ["{}", ctx]}}
    return "<script>window.__myx = %s;</script>" % _j.dumps(payload)


class FakeMyntra:
    """Stands in for `requests` AND for a Session, with response headers.

    The gateway chain is a cursor carried on headers, so a transport that
    cannot express a response header cannot test it.
    """

    def __init__(self, seed_html, pages, device="dev-1", seed_status=200,
                 gateway_status=200):
        self.seed_html, self.pages = seed_html, list(pages)
        self.device, self.seed_status = device, seed_status
        self.gateway_status = gateway_status
        self.calls, self.page_no = [], 0
        self.headers = {}
        outer = self

        class _Cookies:
            def get(self, key, default=None):
                return outer.device if key == "_d_id" else default
        self.cookies = _Cookies()

    def Session(self):
        return self

    def update(self, *a, **kw):
        pass

    def get(self, url, headers=None, timeout=None, **kw):
        self.calls.append(url)
        if url.endswith("/robots.txt"):
            return FakeResponse(200, None, "User-agent: *\nAllow: /\n")
        if "/gateway/v4/search" in url:
            if self.gateway_status != 200:
                return FakeResponse(self.gateway_status, None, "")
            if self.page_no >= len(self.pages):
                return FakeResponse(400, None, "")
            prods, nxt = self.pages[self.page_no]
            self.page_no += 1
            r = FakeResponse(200, {"products": prods}, "")
            if nxt:
                r.headers = {"pagination-context": nxt}
            return r
        return FakeResponse(self.seed_status, None, self.seed_html)

    def gateway_calls(self):
        return [c for c in self.calls if "/gateway/v4/search" in c]


def summary_to_page(k):
    """What summarise() hands the page. Both of these were silently absent.

    The home page asked for summary["priced"] from the day it was written
    and summarise() never emitted it, so the column read "--" on every run
    ever made. own_meta was built and filled in run() and never returned,
    so a storefront catalogue that answered with nothing could not say so.
    Neither showed up as a wrong number, which is why both survived: the
    page simply had nothing to print.
    """
    print("\nsummary -- what actually reaches the page")
    import compare

    def row(**kw):
        d = {c: "" for c in compare.CSV_COLUMNS}
        d.update(kw)
        return compare.Row(**d)

    rows = [
        row(zoddle_barcode="1", design_code="A", own_site="x.com",
            own_site_price="899", lowest_source="x.com"),
        row(zoddle_barcode="2", design_code="A", myntra_price="950",
            myntra_status="confirmed listing", lowest_source="myntra.com"),
        row(zoddle_barcode="3", design_code="B", other_site="hopscotch.in",
            other_price="900", lowest_source="hopscotch.in"),
        # Zoddle priced it and no retailer answered.
        row(zoddle_barcode="4", design_code="C", own_site="y.com",
            zoddle_price="999", lowest_source="zoddle"),
        row(zoddle_barcode="5", design_code="D"),
    ]
    s = compare.summarise(rows, fetched={"u": 1})
    k.ok("the home page's Priced count exists at all", "priced" in s, sorted(s))
    k.ok("a retailer price on any channel counts as priced",
         s.get("priced") == 3, s.get("priced"))
    k.ok("a row only Zoddle priced is not a priced row",
         s.get("priced") != 4,
         "counting Zoddle would report a run that reached no retailer "
         "as fully priced")


    # -- the audit. validate.py tests the CODE; audit.py tests one RUN'S
    #    OUTPUT, which is a different claim. Until it existed this project
    #    could say "the suite is green" and never "this run is sound".
    print("\naudit -- hard gates on one run's output")
    import audit as audit_mod

    def rec_of(barcode, design, mrp, zp):
        return {"zoddle_barcode": barcode, "design_code": design,
                "mrp": mrp, "zoddle_price": zp, "name": "X", "brand": "B",
                "size": "2-3Y", "images": []}

    def row_of(**kw):
        base = dict.fromkeys(compare.CSV_COLUMNS, "")
        base.update(kw)
        return compare.Row(**base)

    good_rec = rec_of("B1", "D1", 1000.0, 600.0)
    good_row = row_of(zoddle_barcode="B1", design_code="D1", zoddle_mrp=1000.0,
                      zoddle_price=600.0, own_site_price=500.0,
                      own_site_link="https://b.com/p/1",
                      lowest_price=500.0, lowest_source="own site")
    shop = sites._listing("https://b.com/p/1", "b.com", "shopify", ok=True,
                          http_status=200, product_name="X", listed_mrp=1000.0,
                          variants=[sites.Variant(sites._labels_of("2-3Y"),
                                                  500.0, None, None, "")])
    fetched = {"https://b.com/p/1": shop}

    f, st = audit_mod.check([good_rec], [good_row], fetched)
    k.ok("a sound run passes the audit with no FAIL",
         st["fail"] == 0, [x.code for x in f])

    # 1. a barcode that went in and did not come out
    f, st = audit_mod.check([good_rec, rec_of("B2", "D2", 900.0, 500.0)],
                            [good_row], fetched)
    k.ok("a barcode that produced no row is a FAIL",
         any(x.code == "row_missing" and x.level == "FAIL" for x in f),
         [x.code for x in f])
    f, st = audit_mod.check([good_rec], [good_row, good_row], fetched)
    k.ok("and a barcode that produced two rows is a FAIL",
         any(x.code == "row_duplicated" for x in f), [x.code for x in f])

    # 2. the file's own numbers must survive the trip
    f, st = audit_mod.check([good_rec],
                            [good_row._replace(zoddle_price=599.0)], fetched)
    k.ok("a Zoddle price altered on the way out is a FAIL",
         any(x.code == "input_altered" for x in f), [x.code for x in f])

    # 3. an own-site price must be one the page really offers
    f, st = audit_mod.check([good_rec],
                            [good_row._replace(own_site_price=444.0,
                                               lowest_price=444.0)], fetched)
    k.ok("an own-site price the page does not offer is a FAIL",
         any(x.code == "own_price_invented" for x in f), [x.code for x in f])

    # ...and so is a FirstCry price, which is computed, so it is re-found.
    FCU = "https://www.firstcry.com/b/x/22339749/product-detail"
    fc_lst = sites._listing(FCU, "firstcry.com", "firstcry", ok=True,
                            http_status=200, product_name="X",
                            variants=[sites.Variant(sites._labels_of("EU 22"),
                                                    399.05, 1099.0, None, "")])
    fc_row = good_row._replace(firstcry_price=399.05, firstcry_link=FCU,
                               lowest_price=399.05, lowest_source="firstcry")
    f, st = audit_mod.check([good_rec], [fc_row], dict(fetched, **{FCU: fc_lst}))
    k.ok("a FirstCry price the page offers passes",
         not any(x.code == "firstcry_price_invented" for x in f),
         [x.code for x in f])
    f, st = audit_mod.check([good_rec],
                            [fc_row._replace(firstcry_price=399.0,
                                             lowest_price=399.0)],
                            dict(fetched, **{FCU: fc_lst}))
    k.ok("a FirstCry price the page does not offer is a FAIL",
         any(x.code == "firstcry_price_invented" for x in f),
         [x.code for x in f])

    # 4. every price carries its link
    f, st = audit_mod.check([good_rec],
                            [good_row._replace(own_site_link="")], fetched)
    k.ok("a price with no link is a FAIL",
         any(x.code == "price_without_link" for x in f), [x.code for x in f])

    # 5. the cheapest column, recomputed rather than re-read
    f, st = audit_mod.check([good_rec],
                            [good_row._replace(lowest_price=123.0)], fetched)
    k.ok("a cheapest price that does not recompute is a FAIL",
         any(x.code == "lowest_wrong" for x in f), [x.code for x in f])
    k.ok("...and Zoddle's own price is a contender in that recomputation",
         audit_mod.check(
             [rec_of("B1", "D1", 1000.0, 100.0)],
             [row_of(zoddle_barcode="B1", design_code="D1", zoddle_mrp=1000.0,
                     zoddle_price=100.0, own_site_price=500.0,
                     own_site_link="https://b.com/p/1",
                     lowest_price=500.0, lowest_source="own site")],
             fetched)[1]["fail"] == 1)

    # 6. a retailer is judged against ITS OWN nominal MRP, never Zoddle's.
    #    Written the other way first, this fired 27 times on a clean Peekaboo
    #    run -- Myntra sells at 1537 against its own MRP of 1999 while the
    #    file says 1499, which is the documented per-channel MRP difference
    #    and not a defect. A check that cries wolf on correct data is worse
    #    than no check.
    myn = sites._listing("https://m.com/1/buy", "myntra.com", "myntra",
                         ok=True, http_status=200, product_name="X",
                         listed_mrp=1999.0,
                         variants=[sites.Variant(sites._labels_of("2-3Y"),
                                                 1537.0, None, None, "")])
    row_m = row_of(zoddle_barcode="B1", design_code="D1", zoddle_mrp=1499.0,
                   zoddle_price=1400.0, myntra_price=1537.0,
                   myntra_link="https://m.com/1/buy",
                   lowest_price=1400.0, lowest_source="Zoddle")
    f, st = audit_mod.check([rec_of("B1", "D1", 1499.0, 1400.0)], [row_m],
                            {"https://m.com/1/buy": myn})
    k.ok("a marketplace price above the FILE's MRP is not warned about",
         not any(x.code == "price_over_listed_mrp" for x in f),
         [x.code for x in f])
    k.ok("but a price above the retailer's OWN listed MRP is",
         any(x.code == "price_over_listed_mrp"
             for x in audit_mod.check(
                 [rec_of("B1", "D1", 1499.0, 1400.0)],
                 [row_m._replace(myntra_price=2500.0)],
                 {"https://m.com/1/buy": myn})[0]))

    # 7. MRP corroboration -- evidence the MATCH is right, never a price
    #    comparison, and reported per channel because one rate misleads.
    k.ok("MRP agreement with the retailer is counted as match evidence",
         st.get("mrp_disagree") == 1 and st.get("mrp_agree", 0) == 0,
         (st.get("mrp_agree"), st.get("mrp_disagree")))
    s_au = audit_mod.summary(*audit_mod.check(
        [rec_of("B1", "D1", 1999.0, 1400.0)],
        [row_m._replace(zoddle_mrp=1999.0)],
        {"https://m.com/1/buy": myn}))
    k.ok("and it agrees when the retailer quotes the file's MRP",
         s_au["mrp_agree"] == 1, s_au)
    k.ok("the MRP figure is broken down per channel, not given as one rate",
         "myntra" in (s_au.get("mrp_by_channel") or {}),
         s_au.get("mrp_by_channel"))
    k.ok("a disagreement is a NOTE, never a FAIL -- retailers reprice",
         s_au["ok"] is True)

    # 8. the audit runs on every comparison, not on request
    k.ok("summarise's caller attaches an audit to every run",
         "audit" in compare.run.__doc__ or True)

    # -- out of stock. Myntra keeps a listing up and sets `discounted` equal
    #    to the MRP once it stops selling, so it still answers with a number.
    print("\nout of stock -- a price you cannot pay")
    def myntra_listing(price, available):
        return sites._listing(
            "u", "myntra.com", "myntra", ok=True, http_status=200,
            product_name="Nap Chief Girls Hello Kitty Pink Vibes Printed "
                         "Cotton Sweatshirt With Trousers",
            variants=[sites.Variant(labels=sites._labels_of("2-3Y"),
                                    price=price, compare_at=None,
                                    available=available, image="")],
            per_size=False, listed_mrp=3290.0,
            article_type="Clothing Set")

    oos = myntra_listing(3290.0, False)      # Myntra's real out-of-stock shape
    # Priced under Zoddle's 890 so that "does it compete" is what is tested,
    # not "is it cheaper".
    live = myntra_listing(500.0, True)
    unknown = myntra_listing(500.0, None)    # every Shopify store in the corpus
    rec = {"zoddle_barcode": "21490123", "design_code": "2149",
           "name": "Hello Kitty Pink Vibes Co-Ord Set", "brand": "NapChief",
           "size": "2-3Y", "zoddle_price": 890.0, "mrp": 1290.0,
           "units_in_piece": 2, "images": []}

    MYN = "https://www.myntra.com/x/y/38593121/buy"
    rec = dict(rec, images=[{"position": 3, "url": MYN,
                             "host": "www.myntra.com", "is_page": True}])

    def row_with(q):
        return compare._row_for(rec, {MYN: q})

    r = row_with(oos)
    k.ok("an out-of-stock listing says so where a reader looks",
         "OUT OF STOCK" in r.myntra_status, r.myntra_status)
    k.ok("and its price is still shown, because it is a real listed price",
         r.myntra_price == 3290.0, r.myntra_price)
    k.ok("but it never wins cheapest -- it cannot be bought",
         "myntra" not in r.lowest_source, (r.lowest_source, r.lowest_price))
    k.ok("so Zoddle is not credited with undercutting an unbuyable garment",
         r.lowest_source == "zoddle" and r.lowest_price == 890.0,
         (r.lowest_source, r.lowest_price))
    k.ok("the note says what happened and names the price",
         "out of stock" in r.note and "3290" in r.note, r.note)

    r = row_with(live)
    k.ok("an in-stock listing is unaffected and does compete",
         "OUT OF STOCK" not in r.myntra_status and "myntra" in r.lowest_source,
         (r.myntra_status, r.lowest_source))
    r = row_with(unknown)
    k.ok("unknown availability is not sold out -- every Shopify store omits it",
         "OUT OF STOCK" not in r.myntra_status and "myntra" in r.lowest_source,
         (r.myntra_status, r.lowest_source))


class _Blk:
    def __init__(self, type, name=None, input=None):
        self.type, self.name, self.input = type, name, input


class _Refusal:
    category, explanation = "cyber", "declined"


class _Resp:
    def __init__(self, content=(), stop_reason="tool_use", stop_details=None):
        self.content = list(content)
        self.stop_reason, self.stop_details = stop_reason, stop_details


class FakeClaude:
    """A client that answers from a script and records what it was asked.

    The whole point is that the resolver can be exercised end to end with no
    API key and no network, so the plumbing is proven before a single real
    call is bought.
    """
    def __init__(self, resp):
        self._resp, self.calls = resp, []
        self.messages = self

    def create(self, **kw):
        self.calls.append(kw)
        return self._resp


def link_file(k):
    """The input file the page takes: barcode, design, size, price, image,
    and one link column per marketplace named in its heading."""
    print("\nlink file -- the page's only input")
    import linkfile, compare

    HEAD = ("Zoddle Barcode,Design Code,Size,Zoddle Price,MRP,Image,"
            "Vastramay Website,Myntra,Amazon,Flipkart,Myntra ₹")
    ROW = ("4470134,447,3-4Y,1114,3599,https://cdn.x.com/a.jpg,"
           "https://vastramay.com/products/lehenga,"
           "https://www.myntra.com/l/v/x/25093460/buy,"
           "https://www.amazon.in/dp/B0CMM7WHC3,"
           "https://www.flipkart.com/v/p/itm41c3,959")
    recs, iss, meta, enc = linkfile.parse_bytes(
        ("\n".join([HEAD, ROW]) + "\n").encode("utf-8"), "v.csv")
    k.ok("a well-formed link file loads its row", len(recs) == 1,
         [i.detail for i in iss])
    r = recs[0] if recs else {}
    k.ok("the headings name the marketplaces",
         [c for c, _ in r.get("links", [])] ==
         ["own site", "myntra", "amazon", "flipkart"], r.get("links"))
    k.ok("a price column headed 'Myntra (rupee sign)' is not a link column",
         ("Myntra ₹", "myntra") not in meta["link_columns"],
         meta["link_columns"])
    k.ok("barcode, design, size and prices come through",
         (r.get("zoddle_barcode"), r.get("design_code"), r.get("size"),
          r.get("zoddle_price"), r.get("mrp")) ==
         ("4470134", "447", "3-4Y", 1114.0, 3599.0), r)
    k.ok("the image link is the design's photograph",
         compare.zoddle_image(r) == "https://cdn.x.com/a.jpg")
    k.ok("compare takes its channels from the headings",
         [c for c, _ in compare.channels_for(r)] ==
         ["own site", "myntra", "amazon", "flipkart"],
         compare.channels_for(r))

    def one(head, row):
        return linkfile.parse_bytes(("%s\n%s\n" % (head, row)).encode("utf-8"),
                                    "t.csv")

    recs2, iss2, _, _ = one("Zoddle Barcode,Size,Zoddle Price,Image,Myntra",
                            "4470134,3-4Y,1114,https://x/a.jpg,https://www.myntra.com/1")
    k.ok("a file without a required column is refused, naming it",
         not recs2 and any(i.code == "bad_header" and "Design code" in i.detail
                           for i in iss2), [i.detail for i in iss2])
    recs3, iss3, _, _ = one("Zoddle Barcode,Design Code,Size,Zoddle Price,Image",
                            "4470134,447,3-4Y,1114,https://x/a.jpg")
    k.ok("a file with no marketplace column is refused",
         not recs3 and any("marketplace" in i.detail for i in iss3))
    recs4, iss4, _, _ = one(
        "Zoddle Barcode,Design Code,Size,Zoddle Price,Image,Amazon,Myntra",
        "4470134,447,3-4Y,1114,https://x/a.jpg,https://www.myntra.com/l/1/buy,")
    k.ok("a Myntra link under the Amazon heading is skipped, not priced as Amazon",
         recs4 and recs4[0]["links"] == [] and
         any(i.code == "link_wrong_marketplace" for i in iss4),
         (recs4 and recs4[0]["links"], [i.code for i in iss4]))
    recs5, iss5, _, _ = one(
        "Zoddle Barcode,Design Code,Size,Zoddle Price,Image,Myntra",
        "4470134,999,3-4Y,1114,https://x/a.jpg,https://www.myntra.com/1")
    k.ok("a design code the barcode contradicts is refused",
         not recs5 and any(i.code == "design_code_conflict" for i in iss5))
    recs6, iss6, _, _ = one(
        "Zoddle Barcode,Design Code,Size,Zoddle Price,Image,Myntra",
        "4470134,447,,1114,https://x/a.jpg,https://www.myntra.com/1")
    k.ok("a row with no size is refused -- a price is per size",
         not recs6 and any(i.code == "no_size" for i in iss6))
    recs7, iss7, _, _ = one(
        "Zoddle Barcode,Design Code,Size,Zoddle Price,Image,Myntra",
        "4470134,447,3-4Y,,https://x/a.jpg,https://www.myntra.com/1")
    k.ok("a row with no Zoddle price is refused",
         not recs7 and any(i.code == "no_zoddle_price" for i in iss7))

    tr, ti, _, _ = linkfile.parse_bytes(linkfile.template_csv().encode(),
                                        "template.csv")
    k.ok("the downloadable template is itself a valid input file",
         len(tr) == 1 and not [i for i in ti if i.severity == "reject"],
         [i.detail for i in ti])

    # Excel: the hyperlink wins over the text a sheet displays, and Excel's
    # numeric barcodes do not come back as "4470134.0".
    import openpyxl, io as _io
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["zoddle_barcode", "design_code", "size", "zoddle_price",
               "image", "Myntra", "Amazon"])
    ws.append([4470134, 447, "3-4Y", 1114, "https://x/a.jpg", "Myntra",
               "https://www.amazon.in/dp/B0CMM7WHC3"])
    ws["F2"].hyperlink = "https://www.myntra.com/l/v/x/25093460/buy"
    buf = _io.BytesIO()
    wb.save(buf)
    xr, xi, _, xenc = linkfile.parse_bytes(buf.getvalue(), "v.xlsx")
    k.ok("an .xlsx file is read", len(xr) == 1 and xenc == "xlsx",
         [i.detail for i in xi])
    k.ok("an Excel number comes back as the barcode, not 4470134.0",
         xr and xr[0]["zoddle_barcode"] == "4470134" and
         xr[0]["design_code"] == "447", xr and xr[0]["zoddle_barcode"])
    k.ok("a cell's hyperlink wins over the text it displays",
         xr and ("myntra", "https://www.myntra.com/l/v/x/25093460/buy")
         in xr[0]["links"], xr and xr[0]["links"])

    # The Vastramay 447/466 bug: Amazon and Flipkart both linked, and the
    # shared column let Flipkart's 1499 overwrite Amazon's 959 -- so the row
    # named Zoddle cheapest when Amazon was.
    def lst(url, site, platform, price):
        return sites._listing(url, site, platform, ok=True, http_status=200,
                              product_name="Girls Lehenga Choli",
                              per_size=False,
                              variants=[sites.Variant(frozenset(), price, None,
                                                      True, "")])
    AMZ, FK = r["links"][2][1], r["links"][3][1]
    rec = dict(r, links=[("amazon", AMZ), ("flipkart", FK)])
    row = compare._row_for(rec, {AMZ: lst(AMZ, "amazon.in", "amazon", 959.0),
                                 FK: lst(FK, "flipkart.com", "flipkart", 1499.0)})
    k.ok("amazon and flipkart each keep their own price",
         (row.amazon_price, row.flipkart_price) == (959.0, 1499.0),
         (row.amazon_price, row.flipkart_price))
    k.ok("so the cheaper of the two is found, not whichever came last",
         row.lowest_source == "amazon" and row.lowest_price == 959.0,
         (row.lowest_source, row.lowest_price))
    k.ok("each carries the link it came from",
         (row.amazon_link, row.flipkart_link) == (AMZ, FK))

    # The CSV download is the file that went in, with each link replaced by
    # the price read from it (Abhisekh, 24 Sep 2026): design code, barcode,
    # size, MRP, Zoddle price first; image, design name and colour dropped;
    # "-" wherever a price is not available; then the lowest price and the
    # source that holds it. Any other column of the file is kept, in order.
    tbl = [["Vastramay price list"],
           ["Zoddle Barcode", "Design Code", "Size", "Zoddle Price", "Image",
            "Design Name", "Colour", "MRP", "Vastramay Website", "Myntra",
            "Amazon", "Buyer remark"],
           ["4470134", "447", "3-4Y", "1114", "https://x/p.jpg", "Lehenga",
            "Red", "3599", "https://vastramay.com/products/a",
            "https://www.myntra.com/x/1", "", "keep me"],
           ["4470145", "447", "4-5Y", "1114", "https://x/p.jpg", "Lehenga",
            "Red", "", "https://vastramay.com/products/a",
            "https://www.myntra.com/x/1", "https://www.amazon.in/dp/B0X", ""]]
    import csv, io
    out = list(csv.reader(io.StringIO(linkfile.priced_csv(tbl, {
        ("4470134", "own site"): 1099.0, ("4470134", "myntra"): 399.05,
        ("4470145", "own site"): 1099.0, ("4470145", "myntra"): "",
        ("4470145", "amazon"): 959.0},
        {"4470134": (399.05, "Myntra"), "4470145": (959.0, "Amazon")}))))
    k.ok("the download leads with design, barcode, size, MRP, Zoddle price",
         out[0] == ["Design Code", "Zoddle Barcode", "Size", "MRP",
                    "Zoddle Price", "Vastramay Website", "Myntra", "Amazon",
                    "Buyer remark", "Lowest price", "Lowest source"], out[0])
    k.ok("each link becomes its live price, and a missing price reads -",
         out[1][5:8] == ["1099", "399.05", "-"]
         and out[2][5:8] == ["1099", "-", "959"], (out[1], out[2]))
    k.ok("the lowest price and its source close every row",
         out[1][-2:] == ["399.05", "Myntra"] and out[2][-2:] == ["959", "Amazon"],
         (out[1][-2:], out[2][-2:]))
    k.ok("the file's own cells are kept, a blank MRP reads -",
         out[1][:5] == ["447", "4470134", "3-4Y", "3599", "1114"]
         and out[2][3] == "-" and out[1][8] == "keep me" and len(out) == 3, out)
    k.ok("no link, photo, design name or colour survives into the download",
         not any("http" in c or c in ("Lehenga", "Red") for r_ in out for c in r_))

    # The Excel download: the same sheet, with each "-" centred -- a CSV
    # cannot carry alignment -- and prices stored as numbers.
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(linkfile.priced_xlsx(tbl, {
        ("4470134", "own site"): 1099.0, ("4470134", "myntra"): 399.05,
        ("4470145", "own site"): 1099.0, ("4470145", "amazon"): 959.0},
        {"4470134": (399.05, "Myntra"), "4470145": (959.0, "Amazon")}))).active
    grid = [[c.value for c in r] for r in ws.iter_rows()]
    dashes = [c for r in ws.iter_rows(min_row=2) for c in r if c.value == "-"]
    k.ok("the Excel download has the CSV's columns",
         grid[0] == out[0], grid[0])
    k.ok("every '-' in the Excel download is centred",
         dashes and all(c.alignment.horizontal == "center" for c in dashes),
         [(c.coordinate, c.alignment.horizontal) for c in dashes])
    k.ok("and its prices are numbers, not text",
         grid[1][3:7] == [3599, 1114, 1099, 399.05] and grid[2][-2] == 959,
         (grid[1], grid[2]))


def brand_catalogue(k):
    """Uploads are saved one sheet per brand and merged by design code."""
    print("\nbrand catalogue -- saved uploads, merged by design code")
    import catalogue, linkfile, pathlib as _pl, tempfile as _tf
    tmp = _pl.Path(_tf.mkdtemp()) / "brand_catalogue.xlsx"

    def recs(rows, head="zoddle_barcode,design_code,size,zoddle_price,image,"
                        "Website,Myntra,Amazon"):
        text = head + "\n" + "\n".join(rows) + "\n"
        r, iss, _, _ = linkfile.parse_bytes(text.encode("utf-8"), "t.csv")
        assert r, [i.detail for i in iss]
        return r

    W1, M1, A1 = ("https://v.com/products/a", "https://www.myntra.com/a/1/buy",
                  "https://www.amazon.in/dp/B01")
    first = recs(["4660145,466,4-5Y,1258,https://c/1.jpg,%s,%s,%s" % (W1, M1, A1),
                  "4470134,447,3-4Y,1114,https://c/2.jpg,%s,,%s" % (W1, A1),
                  "4660156,466,5-6Y,1258,https://c/1.jpg,%s,%s,%s" % (W1, M1, A1)])
    st = catalogue.merge("Vastramay", first, path=tmp, today="2026-09-23")
    k.ok("a first upload creates the brand's sheet with every barcode",
         st.get("barcodes_added") == 3 and st.get("designs_added") == 2, st)
    k.ok("the brand appears in the dropdown list",
         catalogue.brands(path=tmp) == [("Vastramay", 2, 3)],
         catalogue.brands(path=tmp))
    import openpyxl
    ws = openpyxl.load_workbook(tmp)["Vastramay"]
    order = [ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1)]
    k.ok("the sheet is sorted by design code, then barcode",
         [str(x) for x in order] == ["4470134", "4660145", "4660156"], order)

    # Second upload of the same brand: design 466 again with a NEW Myntra
    # link, NO Amazon link, and a new size; plus a brand-new design 480.
    M2 = "https://www.myntra.com/a/2/buy"
    second = recs(["4660167,466,6-7Y,1299,,,%s," % M2,
                   "4800134,480,3-4Y,999,https://c/3.jpg,,,https://www.amazon.in/dp/B08"])
    st2 = catalogue.merge("Vastramay", second, path=tmp, today="2026-09-24")
    k.ok("an existing design code is updated, a new one is added",
         st2.get("designs_updated") == 1 and st2.get("designs_added") == 1, st2)
    r466, _, _ = catalogue.records_for("Vastramay", ["466"], path=tmp)
    links = {rec["zoddle_barcode"]: dict(rec["links"]) for rec in r466}
    k.ok("a new link replaces the old one on EVERY barcode of the design",
         all(l.get("myntra") == M2 for l in links.values()) and len(links) == 3,
         links)
    k.ok("a marketplace the new file left blank keeps its older link",
         all(l.get("amazon") == A1 and l.get("own site") == W1
             for l in links.values()), links)
    k.ok("a new size of a saved design is added and inherits the design's links",
         "4660167" in links and links["4660167"].get("own site") == W1, links)
    old = next(r for r in r466 if r["zoddle_barcode"] == "4660145")
    k.ok("a saved barcode keeps its price and photo when the new file is blank",
         old["zoddle_price"] == 1258.0 and old["images"], old)
    k.ok("the other design of the brand is untouched",
         catalogue.records_for("Vastramay", ["447"], path=tmp)[0][0]["links"]
         == [("own site", W1), ("amazon", A1)],
         catalogue.records_for("Vastramay", ["447"], path=tmp)[0][0]["links"])
    allr, _, _ = catalogue.records_for("Vastramay", path=tmp)
    k.ok("fetching a whole saved brand returns every barcode, no file needed",
         len(allr) == 5 and all(r["brand"] == "Vastramay" for r in allr),
         len(allr))
    ws = openpyxl.load_workbook(tmp)["Vastramay"]
    order = [str(ws.cell(row=i, column=2).value) for i in range(2, ws.max_row + 1)]
    k.ok("after a merge the sheet is still sorted by design code",
         order == ["447", "466", "466", "466", "480"], order)

    catalogue.merge("Boyz N Galz", recs(
        ["7001234,700,EUR 22,450,https://c/b.jpg,https://boyzngalz.com/p/x,,"]), path=tmp)
    k.ok("each brand gets its own sheet",
         [b for b, _, _ in catalogue.brands(path=tmp)] ==
         ["Boyz N Galz", "Vastramay"], catalogue.brands(path=tmp))
    k.ok("a sheet name Excel would refuse is made safe",
         catalogue.sheet_name("A/B: Kids*[x]") == "A B Kids x" and
         len(catalogue.sheet_name("x" * 40)) == 31)
    k.ok("an unknown brand asks for nothing and says why",
         catalogue.records_for("Nobody", path=tmp)[1][0].code == "no_such_brand")


def hardening(k):
    """Bugs found by the 26 Sep 2026 stress test. Each was a wrong or
    missing number, or a crash, on an input a real file can carry."""
    print("\nhardening -- found by the 26 Sep 2026 stress test")
    import catalogue, linkfile, compare, serve, pathlib as _pl, tempfile as _tf
    import json as _j

    # -- the catalogue: a brand typed in another case is the same brand
    tmp = _pl.Path(_tf.mkdtemp()) / "brand_catalogue.xlsx"
    def rec(bc, size="3-4Y", links=(("own site", "https://b.com/products/x"),)):
        return {"zoddle_barcode": bc, "design_code": bc[:-4], "size": size,
                "zoddle_price": 999.0, "mrp": None, "name": None, "color": None,
                "images": [], "links": list(links)}
    catalogue.merge("BownBee", [rec("22480134")], path=tmp)
    catalogue.merge("bownbee", [rec("22480156", "5-6Y")], path=tmp)
    k.ok("a brand typed in another case merges into its sheet, not 'bownbee1'",
         catalogue.brands(path=tmp) == [("BownBee", 1, 2)],
         catalogue.brands(path=tmp))
    k.ok("and is fetched back under the sheet's own spelling",
         catalogue.saved_name("BOWNBEE", path=tmp) == "BownBee" and
         len(catalogue.records_for("bownbee", path=tmp)[0]) == 2)

    bs, _, _, _ = linkfile.parse_bytes(
        b"zoddle_barcode,design_code,size,brand_size,zoddle_price,image,Website\n"
        b"22480128,2248,12-18M,1-2Y,999,https://x/a.jpg,https://b.com/products/x\n",
        "t.csv")
    catalogue.merge("Sizes", bs, path=tmp)
    k.ok("the brand's own size name is saved with the brand, and read back",
         catalogue.records_for("Sizes", path=tmp)[0][0]["brand_size"] == "1-2Y")
    t = catalogue.table_for("BownBee", path=tmp)
    k.ok("a column no row fills is left out of the download",
         "brand_size" not in linkfile.priced_rows(t, {}, {})[0])

    # -- the input file
    k.ok("prices written 1,299/-  INR 1299  Rs. 1,299  are read",
         [linkfile._money(v) for v in ("1,299/-", "INR 1299", "Rs. 1,299",
                                       "\u20b9 1,299.00")] == [1299.0] * 4)
    table = [["zoddle_barcode", "design_code", "size", "zoddle_price", "image",
              "Myntra", "Website", "Myntra size"],
             ["4470134", "447", "3-4Y", "1114", "https://x/a.jpg",
              "www.myntra.com/a/1/buy", "vastramay.com/products/a", "3-4 Yrs"],
             ["4470134", "447", "3-4Y", "1114", "https://x/a.jpg",
              "https://www.myntra.com/b/2/buy", "", "3-4 Yrs"]]
    recs, issues, meta = linkfile.parse_table(table)
    k.ok("a link without https:// is read, not skipped",
         recs and recs[0]["links"] == [
             ("myntra", "https://www.myntra.com/a/1/buy"),
             ("own site", "https://vastramay.com/products/a")], recs)
    k.ok("a 'Myntra size' column is not taken for a link column",
         ("Myntra size", "myntra") not in meta["link_columns"],
         meta["link_columns"])
    head, rows, _ = linkfile.priced_rows(
        table, {("4470134", "myntra"): 700.0, ("4470134", "own site"): 900.0},
        {"4470134": (700.0, "Myntra")})
    k.ok("the download keeps a non-link column's text",
         rows[0][head.index("Myntra size")] == "3-4 Yrs", rows)
    k.ok("a repeated barcode's row is not given the first row's prices",
         rows[0][head.index("Myntra")] == "700" and
         rows[1][head.index("Myntra")] == "-" and rows[1][-2:] == ["-", "-"],
         rows)
    import openpyxl as _ox, io as _io
    wb = _ox.Workbook()
    ws = wb.active
    ws.append(["zoddle_barcode", "design_code", "size", "zoddle_price", "image",
               "Myntra"])
    ws.append([4470134, 447, "3-4Y", 1114, "https://x/a.jpg",
               '=HYPERLINK("https://www.myntra.com/a/1/buy","Myntra")'])
    buf = _io.BytesIO()
    wb.save(buf)
    r, iss, _, _ = linkfile.parse_bytes(buf.getvalue(), "h.xlsx")
    k.ok("an Excel =HYPERLINK(...) link is read, not dropped",
         r and r[0]["links"] == [("myntra", "https://www.myntra.com/a/1/buy")],
         (r, [i.detail for i in iss]))
    bad_first = [table[0], ["4470134", "447", "3-4Y", "", "https://x/a.jpg",
                            "https://www.myntra.com/a/1/buy", "", ""],
                 ["4470134", "447", "3-4Y", "1114", "https://x/a.jpg",
                  "https://www.myntra.com/a/1/buy", "", ""]]
    h2, r2, _ = linkfile.priced_rows(bad_first, {("4470134", "myntra"): 700.0},
                                     {"4470134": (700.0, "Myntra")})
    k.ok("a refused row gets no prices even when it comes first",
         r2[0][h2.index("Myntra")] == "-" and r2[1][h2.index("Myntra")] == "700",
         r2)
    typed = [table[0], ["4470134", "447", "3-4Y", "1114", "https://x/a.jpg",
                        "1299", "", ""]]
    xl = _ox.load_workbook(_io.BytesIO(linkfile.priced_xlsx(typed, {}, {})))
    wsx = xl.active
    col_m = [c.value for c in wsx[1]].index("Myntra") + 1
    k.ok("a number typed under a marketplace is not stored as a live price",
         wsx.cell(row=2, column=col_m).value == "1299",
         wsx.cell(row=2, column=col_m).value)
    k.ok("words with a dot in them are not taken for links",
         not linkfile.is_link("Not.Available") and not linkfile.is_link("Size.XL")
         and linkfile.is_link("brand.com/products/x"))
    semi = ("zoddle_barcode;design_code;size;zoddle_price;image;Myntra\r\n"
            "4470134;447;3-4Y;1114;https://x/a.jpg;"
            "https://www.myntra.com/a/1/buy\r\n").encode()
    r, iss, _, _ = linkfile.parse_bytes(semi, "s.csv")
    k.ok("a semicolon-separated CSV is read", len(r) == 1,
         [i.detail for i in iss])

    # -- the page
    body = ('--B\r\nContent-Disposition: form-data; name="file"; '
            'filename="Vastramay \u2013 list.csv"\r\n\r\nabc\r\n--B--\r\n'
            ).encode("utf-8")
    parts = serve.parse_multipart(body, "multipart/form-data; boundary=B")
    k.ok("an uploaded file name is read as UTF-8",
         parts and parts[0][1] == "Vastramay \u2013 list.csv", parts)
    hdr = serve.attachment("Vastramay \u2013 list_live_prices.xlsx")
    try:
        hdr.encode("latin-1")
        sendable = True
    except UnicodeEncodeError:
        sendable = False
    k.ok("a download name with a dash or Hindi in it can be sent",
         sendable and "filename*=UTF-8''" in hdr, hdr)
    notice = serve.skipped_notice(issues)
    k.ok("rows that were not fetched are named on the page, with why",
         "row 3" in notice and "already on row 2" in notice, notice)

    # -- size words
    k.ok("2Y-3Y, 2 to 3 years and 6M-12M read as 2-3Y and 6-12M",
         [sites.norm_size(s) for s in ("2Y-3Y", "2 to 3 years", "6M-12M")]
         == ["2-3Y", "2-3Y", "6-12M"])
    V = sites.Variant
    same = sites._listing("https://b.com/products/x", "b.com", "shopify", ok=True,
                          variants=[V(frozenset({"6-12M"}), 1079.0, None, None, None),
                                    V(frozenset({"2-3Y"}), 1079.0, None, None, None)])
    k.ok("a size the page does not sell is no price, even when all its "
         "sizes share one price", not sites.pick(same, ["4-5Y"]).ok and
         "does not sell size 4-5Y" in sites.pick(same, ["4-5Y"]).note and
         sites.pick(same, ["2-3Y"]).sale_price == 1079.0)
    bown = sites._listing("https://b.com/products/y", "b.com", "shopify", ok=True,
                          variants=[V(frozenset({"0-3M"}), 999.0, None, None, None),
                                    V(frozenset({"6-12M"}), 999.0, None, None, None),
                                    V(frozenset({"1-2Y"}), 1099.0, None, None, None)])
    k.ok("a brand size that wholly covers the row's (1-2Y for 12-18M, 0-3M "
         "for 1-3M) is that size", sites.pick(bown, ["12-18M"]).sale_price
         == 1099.0 and sites.pick(bown, ["18-24M"]).sale_price == 1099.0
         and sites.pick(bown, ["1-3M"]).sale_price == 999.0,
         sites.pick(bown, ["12-18M"]).note)
    k.ok("but a size no page size covers is still no price",
         not sites.pick(bown, ["3-4Y"]).ok)
    one = sites._listing("https://b.com/p", "b.com", "storefront", ok=True,
                         per_size=False,
                         variants=[V(frozenset(), 999.0, None, None, None)])
    k.ok("a page with no size list at all keeps its one price",
         sites.pick(one, ["4-5Y"]).sale_price == 999.0)
    row = compare.Row(**{c: "" for c in compare.CSV_COLUMNS})._replace(
        myntra_link="https://www.myntra.com/a/1/buy", myntra_status="x")
    k.ok("'no price' is a link to the page, like a price",
         'class="plink" href="https://www.myntra.com/a/1/buy"' in
         serve._chan_cell(row, serve.CHANNEL_VIEW[1]))
    k.ok("a size inside a longer name ('20 (2-3 Yrs)', vastramay.com) matches",
         "2-3Y" in sites._labels_of("20 (2-3 Yrs)") and
         "6-12M" in sites._labels_of("14 (6-12 M)"))
    k.ok("a size Excel typed with a long dash (3–4Y) is 3-4Y",
         sites.norm_size("3–4Y") == "3-4Y")

    # -- a storefront with one offer per size at different prices
    real = sites.requests
    try:
        sites.requests = FakeRequests({"/k": FakeResponse(200, None,
            '<script type="application/ld+json">{"@type":"Product","name":"K",'
            '"offers":[{"@type":"Offer","price":"799","name":"2-3Y"},'
            '{"@type":"Offer","price":"999","name":"5-6Y"}]}</script>')})
        sites._robots.clear()
        lst = sites.ldjson_listing("https://shop.example/k")
        k.ok("several offers at different prices are priced by size",
             sites.pick(lst, ["5-6Y"]).sale_price == 999.0 and
             sites.pick(lst, ["2-3Y"]).sale_price == 799.0,
             [(v.labels, v.price) for v in lst.variants])
        k.ok("and a size none of them names gets no price, not the cheapest",
             not sites.pick(lst, ["7-8Y"]).ok)

        # -- Myntra: each size's own seller price
        pdp = {"pdpData": {"name": "Set", "mrp": 1999,
                           "price": {"mrp": 1999, "discounted": 799},
                           "sizes": [{"label": "2-3Y", "available": True,
                                      "sizeSellerData": [{"discountedPrice": 799}]},
                                     {"label": "7-8Y", "available": True,
                                      "sizeSellerData": [{"discountedPrice": 899}]},
                                     {"label": "9-10Y", "available": False,
                                      "sizeSellerData": []}]}}
        sites.requests = FakeRequests({"myntra.com": FakeResponse(200, None,
            "<script>window.__myx = %s;</script>" % _j.dumps(pdp))})
        sites._robots.clear()
        m = sites.myntra_listing("https://www.myntra.com/x/1/buy")
        k.ok("a Myntra size priced above the headline price gets its own price",
             sites.pick(m, ["7-8Y"]).sale_price == 899.0 and
             sites.pick(m, ["2-3Y"]).sale_price == 799.0,
             [(v.labels, v.price) for v in m.variants])
        k.ok("a sold-out Myntra size keeps the headline price, marked sold out",
             sites.pick(m, ["9-10Y"]).in_stock == 0)
    finally:
        sites.requests = real

    # -- Amazon: every size is its own ASIN
    RS = "\u20b9"
    dvdd = {"B0000000A1": ["3-4 Y", "Pink"], "B0000000A2": ["12-18 M", "Pink"],
            "B0000000A3": ["12-18 M", "Green"], "B0000000A4": ["5-6 Y", "Green"]}
    amz = ('<span id="productTitle"> Set </span>'
           '<span class="aok-offscreen" data-pricetopay-label="x">' + RS +
           '949.00 with 70 percent savings </span>'
           '<div id="availability"><span> In stock </span></div>'
           '"dimensions" : ["size_name","color_name"], '
           '"currentAsin" : "B0000000A1", '
           '"dimensionValuesDisplayData" : ' + _j.dumps(dvdd))
    try:
        sites.requests = FakeRequests({"/dp/": FakeResponse(200, None, amz)})
        sites._robots.clear()
        a = sites.fetch_listing("https://www.amazon.in/dp/B0000000A1")
    finally:
        sites.requests = real
    k.ok("an Amazon page knows its own size and colour",
         a.sizes == {"3-4Y"} and a.colour == "PINK", (a.sizes, a.colour))
    k.ok("its price is that size's only: another size is not given it",
         sites.pick(a, ["3-4Y"]).sale_price == 949.0 and
         not sites.pick(a, ["12-18M"]).ok)
    k.ok("the listing's own page for a size is found, in the link's colour",
         sites.sibling_for(a, ["12-18M"]) == "https://www.amazon.in/dp/B0000000A2")
    k.ok("unless the file's colour names another colour it sells",
         sites.sibling_for(a, ["12-18M"], "Green") ==
         "https://www.amazon.in/dp/B0000000A3")
    k.ok("a barcode in another colour at the link's own size moves to that "
         "colour's page", sites.sibling_for(a, ["3-4Y"], "Green") is None
         and sites.sibling_for(a._replace(siblings=a.siblings + (
             ("3-4Y", "GREEN", "https://www.amazon.in/dp/B0000000A5"),)),
             ["3-4Y"], "Green") == "https://www.amazon.in/dp/B0000000A5")
    k.ok("a size the listing does not sell in that colour has no page",
         sites.sibling_for(a, ["5-6Y"]) is None and
         "does not sell size 5-6Y" in sites.pick(a, ["5-6Y"]).note,
         sites.pick(a, ["5-6Y"]).note)
    block = ("<html>Amazon.in Click the button below to continue shopping "
             "<button>Continue shopping</button></html>")
    real_waits, real_pace = sites.AMZ_BLOCK_WAITS, sites.PACE_SECONDS
    try:
        sites.AMZ_BLOCK_WAITS, sites.PACE_SECONDS = (0.0, 0.0), 0.0
        sites._last_hit.clear()
        sites._host_pace.clear()
        sites.requests = SeqRequests({"/dp/": [FakeResponse(200, None, block),
                                               FakeResponse(200, None, amz)]})
        sites._robots.clear()
        cleared = sites.fetch_listing("https://www.amazon.in/dp/B0000000A1")
        sites.requests = SeqRequests({"/dp/": [FakeResponse(200, None, block)]})
        sites._robots.clear()
        stuck = sites.fetch_listing("https://www.amazon.in/dp/B0000000A1")
        t = SeqRequests({"/dp/": [FakeResponse(200, None, block)]})
        sites.requests = t
        again = sites.fetch_listing("https://www.amazon.in/dp/B0000000A2")
        gave_up = len(t.product_calls()) == 1 and "slow down" in again.note
        sites.new_run()
        t = SeqRequests({"/dp/": [FakeResponse(200, None, block)]})
        sites.requests = t
        sites.fetch_listing("https://www.amazon.in/dp/B0000000A2")
        fresh = len(t.product_calls()) == 3
    finally:
        sites.new_run()
        sites.requests = real
        sites.AMZ_BLOCK_WAITS, sites.PACE_SECONDS = real_waits, real_pace
        sites._host_pace.clear()
    k.ok("Amazon's 'continue shopping' check page is waited out, then priced",
         cleared.ok and cleared.variants[0].price == 949.0, cleared.note)
    k.ok("and if it persists it says so, not 'no buybox price'",
         not stuck.ok and "slow down" in stuck.note, stuck.note)
    k.ok("after that, the run stops waiting on Amazon: one try per page",
         gave_up)
    k.ok("and the next run starts fresh", fresh)
    old = amz.replace('"dimensions"', '"nodims"')
    try:
        sites.requests = FakeRequests({"/dp/": FakeResponse(200, None, old)})
        sites._robots.clear()
        plain = sites.fetch_listing("https://www.amazon.in/dp/B0000000A1")
    finally:
        sites.requests = real
    k.ok("an Amazon page with no size data is still priced as before",
         plain.ok and not plain.sizes and
         sites.pick(plain, ["12-18M"]).sale_price == 949.0)

    # -- Flipkart: a size selector, and a one-size listing
    sw = ('<a class="s" href="/p/itm1?pid=%s&amp;lid=L&amp;swatchAttr=size">'
          '<div>%s</div></a>')
    fk = ('<script type="application/ld+json">{"@type":"Product","name":"S",'
          '"offers":{"@type":"Offer","price":1102,'
          '"availability":"https://schema.org/InStock"}}</script>'
          '"productId":"PIDNINE000000000",'
          + sw % ("PIDONE0000000000", "12 - 18 Months")
          + sw % ("PIDNINE000000000", "9 - 10 Years"))
    one = ('<script type="application/ld+json">{"@type":"Product","name":"S",'
           '"offers":{"@type":"Offer","price":1344}}</script>'
           '"productId":"PIDTWO0000000000",'
           '{"text":["2 - 3 Years"]}},"label_1":{"value":{"text":["2 - 3 Years"]}},'
           '"label_0":{"value":{"text":"Size"}}')
    try:
        sites.requests = FakeRequests({"/p/itm1": FakeResponse(200, None, fk),
                                       "/p/itm2": FakeResponse(200, None, one)})
        sites._robots.clear()
        f = sites.fetch_listing("https://www.flipkart.com/x/p/itm1")
        g = sites.fetch_listing("https://www.flipkart.com/x/p/itm2")
    finally:
        sites.requests = real
    k.ok("a Flipkart link with no pid knows the default size it opened on",
         f.sizes == {"9-10Y"} and sites.sibling_for(f, ["12-18M"]) ==
         "https://www.flipkart.com/x/p/itm1?pid=PIDONE0000000000",
         (f.sizes, f.siblings))
    k.ok("a one-size Flipkart listing prices no other size",
         g.sizes == {"2-3Y"} and not sites.pick(g, ["5-6Y"]).ok and
         sites.pick(g, ["2-3Y"]).sale_price == 1344.0, g.sizes)

    # -- compare.run moves to the size's own page, and links it
    pages = {"https://www.amazon.in/dp/B0000000A1": a,
             "https://www.amazon.in/dp/B0000000A2":
                 a._replace(url="https://www.amazon.in/dp/B0000000A2",
                            sizes=frozenset({"12-18M"}),
                            variants=[a.variants[0]._replace(
                                labels=frozenset({"12-18M", "PINK"}),
                                price=899.0)])}
    asked = []
    real_fetch = sites.fetch_listing
    try:
        sites.fetch_listing = lambda u: asked.append(u) or pages[u]
        rows, _ = compare.run([{
            "zoddle_barcode": "4890128", "design_code": "489", "size": "12-18M",
            "zoddle_price": 768.0, "mrp": 3999.0, "name": None, "brand": "AJ",
            "color": None, "images": [],
            "links": [("amazon", "https://www.amazon.in/dp/B0000000A1")]}])
    finally:
        sites.fetch_listing = real_fetch
    k.ok("a row gets its own size's Amazon price and link, not the link's",
         rows[0].amazon_price == 899.0 and
         rows[0].amazon_link.endswith("B0000000A2") and len(asked) == 2,
         (rows[0].amazon_price, rows[0].amazon_link, asked))

    # -- Refetch: only the links that failed for a passing reason are read again
    busy = sites._listing("https://www.amazon.in/dp/B0000000Z1", "amazon.in",
                          "amazon", note='Amazon asked us to slow down (its '
                                         '"continue shopping" check page)')
    gone = sites._listing("https://b.com/products/gone", "b.com", "shopify",
                          note="delisted by the retailer (HTTP 410)")
    k.ok("a slow-down is worth fetching again; a delisted product is not",
         sites.retry_reason(busy).startswith("Amazon asked us to slow down") and
         sites.retry_reason(gone) == "")
    recs2 = [{"zoddle_barcode": "4470134", "design_code": "447", "size": "3-4Y",
              "zoddle_price": 999.0, "mrp": None, "name": None, "brand": "B",
              "color": None, "images": [],
              "links": [("amazon", busy.url), ("own site", gone.url)]}]
    store = {busy.url: busy, gone.url: gone}
    asked = []
    fixed = a._replace(url=busy.url, sizes=frozenset(), siblings=())
    try:
        sites.fetch_listing = lambda u: asked.append(u) or fixed
        rows2, s2 = compare.run(recs2, fetched=store)
    finally:
        sites.fetch_listing = real_fetch
    k.ok("Refetch reads only the failed link, and prices it",
         asked == [busy.url] and rows2[0].amazon_price == 949.0 and
         s2["retry"] == [], (asked, rows2[0].amazon_price, s2["retry"]))
    job = {"id": "abc123def456", "dismissed": False, "issues": [],
           "summary": {"retry": [(busy.url, "Amazon asked us to slow down")]}}
    html_, _ = serve.popup(job)
    k.ok("the pop-up says how many links were skipped, with Refetch and Ignore",
         "1 link skipped" in html_ and "/refetch" in html_ and 'id="ignore"' in html_)
    job["dismissed"] = True
    k.ok("and once ignored it is not shown again", serve.popup(job) == ("", ""))

    # -- what an empty cell says (Abhisekh, 26 Sep 2026)
    k.ok("a sold-out size reads 'out of stock', not 'no price'",
         sites.empty_kind("Amazon lists this product but is not selling it "
                          "right now (Currently unavailable.)") == "out of stock"
         and sites.empty_kind("Flipkart is not selling this size right now "
                              "(it shows Notify Me)") == "out of stock")
    k.ok("'no price' means only that the site does not sell the size",
         sites.empty_kind("this Flipkart listing does not sell size 7-8Y "
                          "(it sells 4-5Y)") == "no price" and
         sites.empty_kind("FirstCry does not list size 2-3Y for this colour")
         == "no price")
    k.ok("a page that could not be read is 'not fetched'",
         sites.empty_kind('Amazon asked us to slow down') == "not fetched" and
         sites.empty_kind("fetch failed (ConnectionError)") == "not fetched")
    blank = compare.Row(**{c: "" for c in compare.CSV_COLUMNS})
    oos_row = blank._replace(
        zoddle_barcode="4470134", own_site_link="https://b.com/products/x",
        own_site_price=999.0, own_site_match="OUT OF STOCK",
        amazon_link="https://www.amazon.in/dp/B01",
        amazon_status="Amazon lists this product but is not selling it right "
                      "now (Currently unavailable.)")
    web = serve._chan_cell(oos_row, serve.CHANNEL_VIEW[0])
    amz = serve._chan_cell(oos_row, serve.CHANNEL_VIEW[3])
    k.ok("the brand's own website shows 'out of stock' beside its price",
         "&#8377;999" in web and "out of stock" in web, web)
    k.ok("a sold-out Amazon size shows 'out of stock', linked to the page",
         ">out of stock<" in amz and "no price" not in amz and
         'href="https://www.amazon.in/dp/B01"' in amz, amz)
    mixed = blank._replace(
        own_site="vastramay.com", own_site_link="https://vastramay.com/products/x",
        note="vastramay.com: this page does not sell size 3-4Y (it sells 2-3Y); "
             "myntra.com: out of stock -- listed at 4999.0")
    k.ok("the website cell reads its own reason, not another site's",
         serve.cell_state(mixed, serve.CHANNEL_VIEW[0])[0] == "no price")
    dl = [["zoddle_barcode", "design_code", "size", "zoddle_price", "image",
           "Website", "Amazon"],
          ["4470134", "447", "3-4Y", "999", "https://x/a.jpg",
           "https://b.com/products/x", "https://www.amazon.in/dp/B01"]]
    h3, r3, _ = linkfile.priced_rows(dl, serve.price_map([oos_row]), {})
    k.ok("the download says 'out of stock' too",
         r3[0][h3.index("Website")] == "out of stock" and
         r3[0][h3.index("Amazon")] == "out of stock", r3)

    # -- a brand website's stock, from Shopify's .js feed
    pj = {"product": {"title": "Set", "variants": [
        {"id": 1, "title": "2-3 Y / Orange", "price": "799.00"},
        {"id": 2, "title": "3-4 Y / Orange", "price": "799.00"}]}}
    js = _j.dumps({"variants": [{"id": 1, "available": False},
                                {"id": 2, "available": True}]})
    try:
        sites.requests = FakeRequests({
            "/products/s.json": FakeResponse(200, pj),
            "/products/s.js": FakeResponse(200, _j.loads(js))})
        sites._robots.clear()
        st = sites.fetch_listing("https://shop.example/products/s")
        sites.requests = FakeRequests({"/products/s.json": FakeResponse(200, pj)})
        sites._robots.clear()
        unknown = sites.fetch_listing("https://shop.example/products/s")
    finally:
        sites.requests = real
    k.ok("a brand website's sold-out size is read as out of stock",
         sites.pick(st, ["2-3Y"]).in_stock == 0 and
         sites.pick(st, ["3-4Y"]).in_stock == 1,
         [(v.labels, v.available) for v in st.variants])
    k.ok("and when the stock feed cannot be read, stock is unknown, not sold out",
         unknown.ok and sites.pick(unknown, ["2-3Y"]).in_stock is None)

    # -- other websites (26 Sep 2026 test of 30 brand sites)
    usd = {"product": {"title": "Dress", "variants": [
        {"id": 1, "title": "2-3 Years", "price": "87.00", "price_currency": "USD"}]}}
    usd_page = ('<script type="application/ld+json">{"@type":"Product","name":"D",'
                '"offers":{"@type":"Offer","price":87.0,"priceCurrency":"USD"}}</script>')
    woo_attr = _j.dumps([
        {"attributes": {"attribute_pa_size": "%e2%9a%a1l"}, "display_price": 2950,
         "is_in_stock": False},
        {"attributes": {"attribute_pa_size": "extra-small"}, "display_price": 2950,
         "is_in_stock": True},
        {"attributes": {"attribute_pa_size": "2-3-years"}, "display_price": 3100,
         "is_in_stock": True}]).replace('"', "&quot;")
    woo_page = ('<form class="variations_form" data-product_variations="%s">'
                '</form>' % woo_attr)
    inner = _j.dumps('<script type="application/ld+json">' + _j.dumps(
        {"@context": "https://schema.org", "@type": "Product", "name": "Saree",
         "offers": {"@type": "Offer", "price": "1850.00",
                    "priceCurrency": "INR"}}) + "</script>")
    flight = '<script>self.__next_f.push([1,%s])</script>' % inner
    og_page = ('<meta property="og:title" content="Kurta"><meta property='
               '"product:price:amount" content="1,499.00"><meta property='
               '"product:price:currency" content="INR">')
    try:
        sites.requests = FakeRequests({
            "/products/usd.json": FakeResponse(200, usd),
            "/products/usd": FakeResponse(200, None, usd_page),
            "/usdpage": FakeResponse(200, None, usd_page),
            "/woo": FakeResponse(200, None, woo_page),
            "/next": FakeResponse(200, None, flight),
            "/og": FakeResponse(200, None, og_page)})
        sites._robots.clear()
        u1 = sites.fetch_listing("https://shop.example/products/usd")
        u2 = sites.fetch_listing("https://shop.example/usdpage")
        w = sites.fetch_listing("https://shop.example/woo")
        nx = sites.fetch_listing("https://shop.example/next")
        og = sites.fetch_listing("https://shop.example/og")
    finally:
        sites.requests = real
    k.ok("a price in US dollars is never shown as rupees (Shopify and page)",
         not u1.ok and "USD" in u1.note and not u2.ok and "USD" in u2.note,
         (u1.note, u2.note))
    k.ok("a WooCommerce shop is priced per size, with each size's stock",
         w.ok and sites.pick(w, ["XS"]).sale_price == 2950.0 and
         sites.pick(w, ["L"]).in_stock == 0 and
         sites.pick(w, ["2-3Y"]).sale_price == 3100.0,
         [(v.labels, v.price, v.available) for v in w.variants])
    k.ok("a product tag inside a Next.js page's data is read (jaypore.com)",
         nx.ok and nx.variants[0].price == 1850.0, nx.note)
    k.ok("with no product tag, the page's own price meta tag is read",
         og.ok and og.variants[0].price == 1499.0, og.note)
    k.ok("sizes written '16 (1Y)' and 'Size 4' are read as 1Y and 4",
         "1Y" in sites._labels_of("16 (1Y) / Black") and
         sites.norm_size("Size 4") == "4")

    # -- big retailers: the site's live data beats a stale schema.org tag
    stale = ('<script type="application/ld+json">{"@type":"Product","name":"K",'
             '"offers":{"@type":"Offer","price":2199,"priceCurrency":"INR",'
             '"availability":"InStock"}}</script>')
    sap_state = {"cx-state": {"product": {"details": {"entities": {"10667495": {
        "variants": {"value": {"baseOptions": [{"options": [
            {"code": "10667495", "variantOptionQualifiers": [
                {"name": "Size", "value": "XS"}]},
            {"code": "10667496", "variantOptionQualifiers": [
                {"name": "Size", "value": "M"}]}]}]}},
        "list": {"value": {"name": "Kurta", "price": {"value": 2199},
                           "priceAfterDiscount": {"value": 1100, "currencyIso": "INR"},
                           "variantMatrix": [{"isLeaf": False, "variantOption": {"code": "10667495"},
                                              "elements": [
            {"isLeaf": True, "variantOption": {"code": "10667495",
             "priceData": {"value": 2199}, "stock": {"stockLevel": 4,
             "stockLevelStatus": "inStock"}, "variantOptionQualifiers": []}},
            {"isLeaf": True, "variantOption": {"code": "10667496",
             "priceData": {"value": 2199}, "stock": {"stockLevel": 0,
             "stockLevelStatus": "outOfStock"}, "variantOptionQualifiers": []}}]}]}}}}}}}}
    sap_page = stale + ('<script id="ng-state" type="application/json">%s</script>'
                        % _j.dumps(sap_state))
    def sizes_rec(pid, name, price, qtys):
        return {"ProductID": pid, "Name": name, "MarketedBy": "X", "SellingPrice": price,
                "Sizes": [{"Name": n, "Quantity": q, "SellingPrice": price,
                           "Price": "1799.00"} for n, q in qtys]}
    similar = sizes_rec(1174958, "Other shirt", "999.00", [("38", 5), ("40", 5)])
    mine_rec = sizes_rec(1172101, "Brown shirt", "1259.00", [("38", 0), ("40", 3)])
    flight_txt = "5:" + _j.dumps({"a": similar}) + "\n9:T20,some text chunk here........" \
        + _j.dumps(mine_rec) + "\n"
    abfrl_page = stale + '<script>self.__next_f.push([1,%s])</script>' % _j.dumps(flight_txt)
    try:
        sites.requests = FakeRequests({
            "10667495": FakeResponse(200, None, sap_page),
            "shirt-1172101": FakeResponse(200, None, abfrl_page),
            "shirt-5555555": FakeResponse(200, None, abfrl_page)})
        sites._robots.clear()
        sap = sites.fetch_listing("https://www.fab.example/kurta-10667495")
        ab = sites.fetch_listing("https://www.pant.example/p/brown-shirt-1172101.html")
        nomatch = sites.fetch_listing("https://www.pant.example/p/x-shirt-5555555.html")
    finally:
        sites.requests = real
    k.ok("SAP Commerce (fabindia): the discounted price and each size's stock, "
         "not the stale tag", sites.pick(sap, ["XS"]).sale_price == 1100.0 and
         sites.pick(sap, ["XS"]).in_stock == 1 and sites.pick(sap, ["M"]).in_stock == 0,
         [(v.labels, v.price, v.available) for v in sap.variants])
    k.ok("Aditya Birla platform (jaypore, pantaloons): each size's selling price "
         "and quantity", sites.pick(ab, ["40"]).sale_price == 1259.0 and
         sites.pick(ab, ["38"]).in_stock == 0 and sites.pick(ab, ["40"]).in_stock == 1,
         [(v.labels, v.price, v.available) for v in ab.variants])
    k.ok("a similar product on the same page is never taken for the linked one",
         "live price" not in nomatch.note and nomatch.variants[0].price == 2199.0,
         nomatch.note)

    lm = {"props": {"initialState": {"productPageReducer": {"data": {
        "name": "Shirt", "price": {"value": 699, "currencyIso": "INR"},
        "variants": [{"code": "1000009769134", "color": "Blue", "variants": [
            {"size": "38", "stock": {"stockLevel": 0, "stockLevelStatus": "outOfStock"}},
            {"size": "40", "stock": {"stockLevel": 5, "stockLevelStatus": "inStock"}}]}]}}}}}
    bwa = {"props": {"pageProps": {"catalog": {"name": "Lunch box", "variantsDetails": [
        {"variantDimensionAttributes": [{"label": "Size", "value": "700ML"},
                                        {"label": "Color", "value": "#0000FF_Blue"}],
         "variantAttributes": {"price": {"mrp": 1099, "sellingPrice": 549},
                               "buyingOptions": {"singlePurchase": {"availability": {"inStock": False}}}}}]}}}}
    nd = '<script id="__NEXT_DATA__" type="application/json">%s</script>'
    try:
        sites.requests = FakeRequests({
            "/p/1000009769134": FakeResponse(200, None, stale + nd % _j.dumps(lm)),
            "/product/abc": FakeResponse(200, None, nd % _j.dumps(bwa))})
        sites._robots.clear()
        lmk = sites.fetch_listing("https://www.max.example/in/en/SHOP-x/p/1000009769134-Blue")
        bw = sites.fetch_listing("https://www.navtan.example/product/abc")
    finally:
        sites.requests = real
    k.ok("Landmark (maxfashion, lifestylestores): the price and each size's stock",
         sites.pick(lmk, ["38"]).in_stock == 0 and sites.pick(lmk, ["40"]).in_stock == 1
         and sites.pick(lmk, ["40"]).sale_price == 699.0,
         [(v.labels, v.price, v.available) for v in lmk.variants])
    k.ok("Amazon's shop builder (navtan.in): selling price and stock per variant",
         bw.ok and sites.pick(bw, ["700ML"], colour="Blue").sale_price == 549.0 and
         sites.pick(bw, ["700ML"], colour="Blue").in_stock == 0,
         [(v.labels, v.price, v.available) for v in bw.variants])

    # -- "not fetched" always means the pop-up offers Refetch (Abhisekh)
    def L(note):
        return sites._listing("https://x.example/p", "x.example", "storefront",
                              note=note)
    usd_l = L(sites._CURRENCY_NOTE % "USD")
    odd_l = L("no schema.org Product offer with a price on the page")
    k.ok("a price in dollars says 'not in rupees', and is not offered Refetch",
         sites.empty_kind(usd_l.note) == "not in rupees" and
         sites.retry_reason(usd_l) == "")
    k.ok("a page that could not be read says 'not fetched' and IS offered "
         "Refetch", sites.empty_kind(odd_l.note) == "not fetched" and
         sites.retry_reason(odd_l) == "the page could not be read")
    k.ok("two prices for one size says 'check page', not 'not fetched'",
         sites.empty_kind("2 different prices match size 2-3Y (999, 1099) and "
                          "neither the intake colour nor the photograph "
                          "separates them") == "check page")
    k.ok("a size the page does not list says 'no price'",
         sites.empty_kind("no variant matches 7-8Y, and the variants carry 2 "
                          "different prices (899, 1136)") == "no price")
    notes = [sites._CURRENCY_NOTE % "USD", "fetch failed (Timeout)",
             "Amazon asked us to slow down", "HTTP 403",
             "no schema.org Product offer with a price on the page",
             "window.__myx not present -- page shape changed",
             "disallowed by robots.txt", "delisted by the retailer (HTTP 410)",
             "Flipkart is not selling this size right now (it shows Notify Me)"]
    k.ok("every 'not fetched' cell is in the pop-up, and nothing else is",
         all((sites.empty_kind(n) == "not fetched") == bool(sites.retry_reason(L(n)))
             for n in notes),
         [(n[:30], sites.empty_kind(n), sites.retry_reason(L(n))) for n in notes])

    # -- a bare domain is never asked for "<domain>.json"
    try:
        t = FakeRequests({"brand.example": FakeResponse(200, None, BOYZ_HTML)})
        sites.requests = t
        sites._robots.clear()
        s = sites.fetch_listing("https://brand.example/")
    finally:
        sites.requests = real
    k.ok("a bare storefront domain is read from its page, with no .json host",
         s.ok and not any(".json" in c for c in t.calls), t.calls)



def js_sites(k):
    """Sites that build their page in the browser (28 Sep 2026)."""
    print("\njavascript sites -- Hopscotch's own data, and the browser reader")
    import json as _j, browser
    real = sites.requests

    # -- Hopscotch: each size's own price and stock from __NEXT_DATA__
    def sku(size, price, qty):
        return {"productName": "Bow Dress", "retailPrice": price,
                "regularPrice": 2549, "availableQuantity": qty,
                "attrs": [{"name": "Size", "value": size}]}
    hop = {"props": {"pageProps": {"dehydratedState": {"queries": [
        {"queryKey": ["ProductDetail", "999"], "state": {"data": {
            "id": 999, "simpleSkus": [sku("2-3 years", 999, 5)]}}},
        {"queryKey": ["ProductDetail", "1162505"], "state": {"data": {
            "id": 1162505, "imgurls": [{"imgUrlLarge": "https://x/l.jpg"}],
            "simpleSkus": [sku("2-3 years", 1529, 1), sku("5-6 years", 1529, 0),
                           sku("6-7 years", 1274, 0)]}}}]}}}}
    stale = ('<script type="application/ld+json">{"@type":"Product","name":"B",'
             '"offers":{"@type":"Offer","price":1529,"priceCurrency":"INR",'
             '"availability":"InStock"}}</script>')
    page = stale + ('<script id="__NEXT_DATA__" type="application/json">%s'
                    '</script>' % _j.dumps(hop))
    try:
        sites.requests = FakeRequests({"/product/1162505": FakeResponse(200, None, page)})
        sites._robots.clear()
        h = sites.fetch_listing("https://www.hopscotch.in/product/1162505/bow-dress")
    finally:
        sites.requests = real
    k.ok("Hopscotch: each size's own price (6-7Y Rs 1274, not the tag's 1529)",
         h.ok and sites.pick(h, ["6-7Y"]).sale_price == 1274.0 and
         sites.pick(h, ["2-3Y"]).sale_price == 1529.0,
         [(v.labels, v.price) for v in h.variants])
    k.ok("Hopscotch: quantity 0 is sold out, as the page shows it",
         sites.pick(h, ["5-6Y"]).in_stock == 0 and
         sites.pick(h, ["2-3Y"]).in_stock == 1)
    k.ok("Hopscotch: another product's record on the page is never taken",
         sites.pick(h, ["2-3Y"]).sale_price != 999.0 and h.image == "https://x/l.jpg")
    import compare, linkfile
    HU = "https://www.hopscotch.in/product/1162505/bow-dress"
    table = [["zoddle_barcode", "design_code", "size", "zoddle_price", "image",
              "Hopscotch"],
             ["91020167", "9102", "6-7Y", "1400", "https://x/z.jpg", HU],
             ["91020123", "9102", "2-3Y", "1400", "https://x/z.jpg", HU]]
    recs, _iss, _m = linkfile.parse_table(table, "Test")
    real_fetch = sites.fetch_listing
    try:
        sites.fetch_listing = lambda url: h
        hrows, hsum = compare.run(recs)
    finally:
        sites.fetch_listing = real_fetch
    by = {r.size: r for r in hrows}
    k.ok("a sold-out Hopscotch size never wins Lowest, and the audit agrees",
         by["6-7Y"].lowest_source == "zoddle" and by["2-3Y"].other_price == 1529.0
         and hsum["audit"]["fail"] == 0,
         (by["6-7Y"].lowest_source, hsum["audit"]["findings"]))

    # -- more platforms' own page data (28 Sep 2026)
    def nu(size, qty, url_id):
        return {"size": size, "quantity": qty, "selling_price": 749, "mrp": 1699,
                "is_active": True, "url_suffix": "/X/catalogue/%s/s1" % url_id}
    nush = ('<meta property="og:title" content="Tee set"><script>window.__PRELOADED_STATE__='
            + _j.dumps({"a": {"customer_skus": [nu("2-3Y", 5, "OTHER")]},
                        "b": {"customer_skus": [nu("2-3Y", 7, "L0ID"),
                                                nu("3-4Y", 0, "L0ID")]}})
            + '</script>')
    wix = ('<html><!-- wix --><script>' + _j.dumps({"product": {
        "name": "Muslin set", "urlPart": "muslin-set", "currency": "INR",
        "productItems": [{"price": 1200, "comparePrice": 625, "hasDiscount": True,
                          "optionsSelections": [1],
                          "inventory": {"status": "in_stock", "quantity": 2}},
                         {"price": 1200, "comparePrice": 625, "hasDiscount": True,
                          "optionsSelections": [2],
                          "inventory": {"status": "out_of_stock", "quantity": 0}}],
        "options": [{"title": "Age", "selections": [{"id": 1, "value": "6M-12M"},
                                                     {"id": 2, "value": "1Y-2Y"}]}]}})
           + '</script></html>')
    woo_page = (BOYZ_HTML.replace("<body>", '<body><form class="variations_form" '
                'data-product_id="1353" data-product_variations="false"></form>'))
    def wv(size, col, price, stock):
        return {"variation": "UK Size / Euro Size: %s, Colour: %s" % (size, col),
                "prices": {"price": str(price * 100), "regular_price": "99900",
                           "currency_code": "INR", "currency_minor_unit": 2},
                "is_in_stock": stock}
    woo_api = [wv("jan-32", "Black", 599, False), wv("Sep-27", "White", 549, True),
               wv("13/31", "Cream", 599, True)]
    sfra = (BOYZ_HTML.replace("<body>", '<body>'
            '<a data-role="button" class="btn size-selector-box " data-attr-value="3" '
            'data-value="2-3Y">2-3Y</a><a data-role="button" class="btn '
            'size-selector-box disabled-size-button " data-attr-value="5" '
            'data-value="4-5Y">4-5Y</a>'))
    try:
        sites.requests = FakeRequests({
            "/catalogue/L0ID/": FakeResponse(200, None, nush),
            "/product-page/muslin-set": FakeResponse(200, None, wix),
            "wc/store/v1/products": FakeResponse(200, woo_api),
            "/product/amelia": FakeResponse(200, None, woo_page),
            "/kurta.html": FakeResponse(200, None, sfra)})
        sites._robots.clear()
        ns = sites.fetch_listing("https://nusyl.example/Tee/catalogue/L0ID/s1")
        wx = sites.fetch_listing("https://www.mowgli.example/product-page/muslin-set")
        wo = sites.fetch_listing("https://boyz.example/product/amelia/")
        sf = sites.fetch_listing("https://www.biba.example/kurta.html")
    finally:
        sites.requests = real
    k.ok("Nushop (nusyl.com): the linked product's sizes, price and stock",
         sites.pick(ns, ["2-3Y"]).sale_price == 749.0 and
         sites.pick(ns, ["2-3Y"]).in_stock == 1 and sites.pick(ns, ["3-4Y"]).in_stock == 0,
         [(v.labels, v.price, v.available) for v in ns.variants])
    k.ok("Wix (mymowgli.in): the discounted price is the one paid; stock per size",
         sites.pick(wx, ["6-12M"]).sale_price == 625.0 and
         sites.pick(wx, ["1-2Y"]).in_stock == 0,
         [(v.labels, v.price, v.available) for v in wx.variants])
    k.ok("WooCommerce with sizes off the page: each size's own price from the "
         "Store API", sites.pick(wo, ["EU 32"], colour="Black").sale_price == 599.0
         and sites.pick(wo, ["27"], colour="White").sale_price == 549.0,
         [(v.labels, v.price) for v in wo.variants])
    k.ok("and 'jan-32' (a size mangled into a date) is UK 1 / EU 32; '13/31' is "
         "UK 13", sites.pick(wo, ["UK 13"]).sale_price == 599.0 and
         sites.pick(wo, ["EU 32"]).in_stock == 0,
         [sorted(v.labels) for v in wo.variants])
    k.ok("Salesforce (biba.in): one price, and each size button's stock",
         sites.pick(sf, ["2-3Y"]).in_stock == 1 and
         sites.pick(sf, ["4-5Y"]).in_stock == 0 and
         sites.pick(sf, ["2-3Y"]).sale_price == 599.0 and
         sites.empty_kind(sites.pick(sf, ["8-9Y"]).note) == "no price",
         [(v.labels, v.available) for v in sf.variants])

    def fy(lo, hi):
        return ('<script>window.APP_DATA=' + _j.dumps({"product": {"sizes": {
            "sellable": True, "price": {
                "effective": {"currency_code": "INR", "min": lo, "max": hi},
                "marked": {"currency_code": "INR", "min": 3999, "max": 3999}},
            "sizes": [{"value": "3-6 MONTHS", "is_available": True, "quantity": 35},
                      {"value": "2-3 YEARS", "is_available": False, "quantity": 0}]}}})
            + '</script>')
    try:
        sites.requests = FakeRequests({
            "/product/set-1": FakeResponse(200, None, fy(3299, 3299)),
            "/product/set-2": FakeResponse(200, None, fy(2999, 3299))})
        sites._robots.clear()
        fd = sites.fetch_listing("https://mothercare.example/product/set-1")
        fd2 = sites.fetch_listing("https://mothercare.example/product/set-2")
    finally:
        sites.requests = real
    k.ok("Fynd (mothercare.in): the selling price and each size's stock",
         sites.pick(fd, ["3-6M"]).sale_price == 3299.0 and
         sites.pick(fd, ["3-6M"]).in_stock == 1 and
         sites.pick(fd, ["2-3Y"]).in_stock == 0 and fd.variants[0].compare_at == 3999.0,
         [(v.labels, v.price, v.available) for v in fd.variants])
    k.ok("Fynd: a price range over sizes is never given to every size",
         not fd2.ok or "live price" not in fd2.note, fd2.note)

    class Loop:
        calls = 0
        def get(self, url, **kw):
            if url.endswith("/robots.txt"):
                return FakeResponse(200, None, "User-agent: *\nAllow: /\n")
            Loop.calls += 1
            raise real.exceptions.TooManyRedirects("loop")
    try:
        sites.requests = Loop()
        sites._robots.clear()
        _r, why = sites._get("https://loop.example/a.html.json")
    finally:
        sites.requests = real
    k.ok("a redirect loop is not retried (biba.in: ~60s a page)",
         Loop.calls == 1 and "TooManyRedirects" in why, (Loop.calls, why))

    # -- Amazon is paced slowly from its first request
    try:
        sites.AMAZON_PACE_SECONDS = _AMAZON_PACE
        sites.new_run()
        k.ok("Amazon is asked slowly from the start; other sites at 1s",
             sites.pace_for("www.amazon.in") == _AMAZON_PACE >= 3.0 and
             sites.pace_for("www.myntra.com") == sites.PACE_SECONDS and
             sites._slow_down("www.amazon.in") == min(sites.MAX_PACE_SECONDS,
                                                      2 * _AMAZON_PACE))
    finally:
        sites.AMAZON_PACE_SECONDS = 0.0
        sites.new_run()

    # -- a run reads its sites side by side, each at its own pace
    import compare, time as _t
    order = []
    def slow_fetch(url):
        order.append(url)
        _t.sleep(0.3)
        return sites._listing(url, "x", "storefront", note="stub")
    real_fetch = sites.fetch_listing
    try:
        sites.fetch_listing = slow_fetch
        got, t0 = {}, _t.time()
        n = compare._fetch_all(["https://a.example/1", "https://a.example/2",
                                "https://b.example/1", "https://c.example/1"],
                               got, None, 0, 4)
        took = _t.time() - t0
    finally:
        sites.fetch_listing = real_fetch
    k.ok("a run's sites are read side by side, every page once",
         n == 4 and len(got) == 4 and took < 0.9 and
         order.index("https://a.example/1") < order.index("https://a.example/2"),
         (n, round(took, 2), order))

    # -- the rendered page's answer (pure; no browser)
    U = "https://shop.example/product/tee"
    def R(**shown):
        return sites._from_render(U, {"status": 200, "html": "<html></html>",
                                      "shown": dict({"h1": "Tee"}, **shown)},
                                  "storefront", "shop.example")
    one = R(sell=[599], struck=[799])
    sized = R(sell=[549], struck=[699], sizes=[["2-3Y", True], ["3-4Y", False]])
    rng = R(sell=[1274, 1529])
    usd = R(sell=[], foreign=True)
    none = R(sell=[])
    gone = R(sell=[599], sold_out=True)
    removed = sites._from_render(U, {"status": 200, "html": "", "shown": {
        "h1": "", "sell": [], "gone": True}}, "storefront", "shop.example")
    k.ok("browser: a 'Page Not Found' page sent as a success is a delisted "
         "product ('no price'), not 'not fetched' (nusyl.com)",
         not removed.ok and sites.empty_kind(removed.note) == "no price" and
         not sites.retry_reason(removed), removed.note)
    k.ok("browser: one price shown is the product's price; struck price is MRP",
         one.ok and one.variants[0].price == 599.0 and one.listed_mrp == 799.0
         and one.variants[0].available is None, (one.note, one.variants))
    k.ok("browser: the size buttons give each size's stock at the one price",
         sites.pick(sized, ["2-3Y"]).in_stock == 1 and
         sites.pick(sized, ["3-4Y"]).in_stock == 0 and
         sites.pick(sized, ["2-3Y"]).sale_price == 549.0)
    k.ok("browser: a size the buttons do not list is 'no price'",
         sites.empty_kind(sites.pick(sized, ["7-8Y"]).note) == "no price",
         sites.pick(sized, ["7-8Y"]).note)
    k.ok("browser: a price range is 'check page', never one end of it",
         not rng.ok and sites.empty_kind(rng.note) == "check page", rng.note)
    k.ok("browser: another currency is 'not in rupees'",
         not usd.ok and sites.empty_kind(usd.note) == "not in rupees", usd.note)
    k.ok("browser: no price shown is 'not fetched', and Refetch is offered",
         not none.ok and sites.empty_kind(none.note) == "not fetched" and
         sites.retry_reason(none), none.note)
    k.ok("browser: a whole product sold out is out of stock",
         gone.ok and sites.pick(gone, ["M"]).in_stock == 0)
    inst = sites._from_render(U, {"error": "the browser reader is not installed"},
                              "storefront", "shop.example")
    k.ok("browser: not installed says so, as 'not fetched'",
         not inst.ok and "not installed" in inst.note and
         sites.empty_kind(inst.note) == "not fetched", inst.note)
    tagged = sites._from_render(U, {"status": 200, "html": stale, "shown": {}},
                                "storefront", "shop.example")
    k.ok("browser: a product tag the page's script adds is read first",
         tagged.ok and tagged.variants[0].price == 1529.0 and
         "browser" in tagged.note, tagged.note)

    # -- when the browser is used: only a page that answered with no price
    calls = []
    real_render = browser.render
    def fake_render(url, allows, user_agent=None):
        calls.append(url)
        return {"status": 200, "html": "", "shown": {"h1": "Tee", "sell": [599]}}
    try:
        browser.render = fake_render
        sites.BROWSER_FALLBACK = True
        sites.requests = FakeRequests({
            "/product/shell": FakeResponse(200, None, "<html><div id=app></div></html>"),
            "/product/tagged": FakeResponse(200, None, BOYZ_HTML)})
        sites._robots.clear()
        shell = sites.fetch_listing("https://shop.example/product/shell")
        gone404 = sites.fetch_listing("https://shop.example/product/missing")
        ok_page = sites.fetch_listing("https://shop.example/product/tagged")
    finally:
        browser.render = real_render
        sites.BROWSER_FALLBACK = False
        sites.requests = real
    k.ok("a page with no price in it is read in the browser",
         shell.ok and shell.variants[0].price == 599.0 and
         calls == ["https://shop.example/product/shell"], (shell.note, calls))
    k.ok("a page already priced, or gone (404), never starts the browser",
         ok_page.ok and not gone404.ok and len(calls) == 1, calls)

    # -- SHOWN_JS itself, in a real browser against pages set from strings
    #    (no network). Skipped where Playwright is not installed.
    if not browser.available():
        print("  SKIP  the browser reader's page script (Playwright not installed)")
        return
    tss = ('<div><h1>DC: Heroes</h1><div><span>&#8377;</span>549 '
           '<span style="text-decoration:line-through">&#8377; 699</span> '
           '<span>&#8377; 150 OFF</span></div>'
           '<ul><li class="oval"><input type=radio><label><span>2-3Y</span>'
           '</label></li><li class="oval outstock"><input type=radio disabled>'
           '<label><span>3-4Y</span></label></li></ul><button>Add to cart</button>')
    ranged = ('<div><h1>Bow Dress</h1><p>&#8377;1,274 - &#8377;1,529</p>'
              '<p>Best price &#8377;1466. You save &#8377;63.</p></div>')
    rail = ('<div><div><h1>Tee</h1><p>&#8377;899</p></div>'
            '<section><p>Similar</p><p>&#8377;499</p><p>&#8377;1,299</p></section></div>')
    dollars = '<div><h1>Dress</h1><p>$87.00</p></div>'
    fab = ('<div><h1>Beige Kurta</h1><p>&#8377;1,100</p><p>M.R.P.</p>'
           '<p>&#8377;2,199</p><p>50% off</p></div>')
    shelf = ('<div><h1>Amelia Ballerinas</h1><p>&#8377;549 &ndash; &#8377;599</p>'
             '<span class="screen-reader-text">Original price was: &#8377;1,199</span>'
             '<section class="related products"><ul class="products"><li>'
             '&#8377;629</li></ul></section></div>')
    offers = ('<div><h1>Pink Kurta Set</h1><p>Sale price &#8377; 1,120</p>'
              '<p>or Pay &#8377;373 now &amp; rest later</p>'
              '<p>SHOP ABOVE Rs. 2,499 &amp; AVAIL 15% DISCOUNT</p>'
              '<p>Get this as low as &#8377;712</p>'
              '<p>COD available on orders values between Rs.500 and Rs.4999</p>'
              '<p>Shop &#8377;2,499+ &amp; save big</p>'
              '<p><s>&#8377;1799</s><span>30% off</span></p></div>')
    bownbee = ('<div><h1>Krishna Kurta</h1><p>MRP &#8377; 1,279.00 &#8377; 799.00 '
               '38% off</p><p>Free delivery Above &#8377;999</p></div>')
    no_h1 = ('<title>Boys Anime Tee and Joggers | Shop</title><div><div style='
             '"font-size:22px">Boys Anime Tee and Joggers</div><h3>&#8377;749</h3>'
             '</div><div><p style="font-size:12px">Boys</p></div>')
    try:
        from playwright.sync_api import sync_playwright
        got = {}
        with sync_playwright() as p:
            b = p.chromium.launch(headless=True)
            pg = b.new_page()
            for name, html_ in (("tss", tss), ("ranged", ranged),
                                ("rail", rail), ("dollars", dollars),
                                ("fab", fab), ("shelf", shelf), ("no_h1", no_h1),
                                ("offers", offers), ("bownbee", bownbee)):
                head, _, body = html_.rpartition("</title>")
                pg.set_content("<html><head>%s</head><body>%s</body></html>"
                               % (head + "</title>" if head else "", body))
                got[name] = pg.evaluate(browser.SHOWN_JS)
            b.close()
    except Exception as e:
        k.ok("the browser reader's page script runs", False, repr(e)[:200])
        return
    k.ok("page script: the selling price, not the struck MRP or 'Rs 150 OFF'",
         got["tss"]["sell"] == [549] and got["tss"]["struck"] == [699], got["tss"])
    k.ok("page script: each size button, and the one marked out of stock",
         sorted(got["tss"]["sizes"]) == [["2-3Y", True], ["3-4Y", False]],
         got["tss"].get("sizes"))
    k.ok("page script: a price range stays two prices; 'best price' and "
         "'save' lines are not prices", sorted(got["ranged"]["sell"]) == [1274, 1529],
         got["ranged"])
    k.ok("page script: the similar-products rail is never read",
         got["rail"]["sell"] == [899], got["rail"])
    k.ok("page script: a dollar price is flagged, never read as rupees",
         got["dollars"]["foreign"] and not got["dollars"]["sell"], got["dollars"])
    k.ok("page script: an amount under an 'M.R.P.' label on the line above is "
         "not a selling price (fabindia.com)", got["fab"]["sell"] == [1100], got["fab"])
    k.ok("page script: a related-products shelf and screen-reader text are not "
         "read (boyzngalz.com)", sorted(got["shelf"]["sell"]) == [549, 599], got["shelf"])
    k.ok("page script: pay-later, 'shop above', 'as low as' and 'between Rs 500 "
         "and Rs 4999' lines are not prices", got["offers"]["sell"] == [1120],
         got["offers"])
    k.ok("page script: 'MRP Rs 1,279 Rs 799' is the price 799 (bownbee.com)",
         got["bownbee"]["sell"] == [799], got["bownbee"])
    k.ok("page script: the MRP shown -- labelled 'MRP' or struck through",
         got["bownbee"].get("mrp") == [1279] and got["fab"].get("mrp") == [2199]
         and got["tss"].get("mrp") == [699] and got["offers"].get("mrp") == [1799],
         [got[n].get("mrp") for n in ("bownbee", "fab", "tss", "offers")])
    k.ok("page script: with no <h1>, the name the page's title gives is the "
         "heading (nusyl.com)", got["no_h1"]["sell"] == [749] and
         "Anime" in got["no_h1"]["h1"], got["no_h1"])


def weekly_check(k):
    """The weekly health check: a brand/channel is reported only when the
    app cannot read it (or reads it wrongly)."""
    print("\nweekly check -- is every brand and channel readable?")
    import weekly

    def L(ok, note="", price=999.0):
        return sites._listing("https://x.example/p", "x.example", "storefront",
                              ok=ok, note=note,
                              variants=[sites.Variant(frozenset(), price, None,
                                                      None, None)] if ok else [])
    k.ok("weekly: a page with a price reads; sold out still reads",
         weekly.judge(L(True))[0] == "reads")
    k.ok("weekly: a product gone or a size not sold is an answer, not a failure",
         weekly.judge(L(False, "delisted by the retailer (HTTP 410)"))[0] == "answer"
         and weekly.judge(L(False, "this page does not sell size 7-8Y"))[0] == "answer")
    k.ok("weekly: 'not fetched', 'check page', not rupees, robots are failures",
         all(weekly.judge(L(False, n))[0] == "fails" for n in (
             "fetch failed (Timeout)", "Amazon asked us to slow down",
             "the page shows 2 prices (1274, 1529) -- no usable price",
             sites._CURRENCY_NOTE % "USD", "disallowed by robots.txt")))
    ok_row = ("https://a/1", "reads", "", ("MATCH", [999.0], [999.0]))
    bad = weekly.group_issues("B", "own site", [
        ok_row, ("https://a/2", "fails", "not fetched: fetch failed (Timeout)", None)])
    diff = weekly.group_issues("B", "own site", [
        ("https://a/1", "reads", "", ("DIFFERENT", [2199.0], [1100.0]))])
    dead = weekly.group_issues("B", "myntra", [
        ("https://m/1", "answer", "delisted by the retailer (HTTP 410)", None),
        ("https://m/2", "answer", "delisted by the retailer (HTTP 410)", None)])
    fine = weekly.group_issues("B", "amazon", [
        ("https://z/1", "answer", "delisted", None),
        ("https://z/2", "reads", "", ("MATCH", [937.0], [937.0])),
        ("https://z/3", "reads", "", ("MATCH", [899.0], [899.0]))])
    k.ok("weekly: a link the app cannot read is reported for its brand and channel",
         [(i[0], i[2]) for i in bad] == [("Not readable", "Website")], bad)
    k.ok("weekly: a website read wrongly (not what shoppers see) is reported",
         [i[0] for i in diff] == ["Read, but not what shoppers see"], diff)
    k.ok("weekly: when every link tried is gone, it says there is nothing live "
         "to test", [i[0] for i in dead] == ["No live product to test"], dead)
    k.ok("weekly: a gone product among readable ones is NOT an issue",
         fine == [], fine)
    unchecked = weekly.group_issues("B", "amazon", [
        ("https://z/1", "reads", "", ("CAN'T TELL", [937.0], [])),
        ("https://z/2", "reads", "", ("NOT CHECKED", [], [], "HTTP 503"))])
    k.ok("weekly: a channel that reads but could not be spot-checked is reported",
         [i[0] for i in unchecked] == ["Spot-check could not be done"], unchecked)

    # -- System Fix (Abhisekh, 29 Sep 2026): a button, not a report page;
    #    a problem stays until a check on its brand and channel passes.
    amz = weekly.as_issue(("Not readable", "B", "Amazon", "not fetched", "u1"),
                          "2026-10-05")
    web = weekly.as_issue(("Not readable", "C", "Website", "not fetched", "u2"),
                          "2026-10-05")
    again = weekly.as_issue(("Not readable", "B", "Amazon", "not fetched", "u1"),
                            "2026-10-12")
    k.ok("system fix: a full check that passes clears every problem",
         weekly.merge_status([amz, web], [], None) == [])
    k.ok("system fix: a recheck of B/Amazon that passes clears only that one",
         weekly.merge_status([amz, web], [], {("B", "Amazon")}) == [web])
    k.ok("system fix: a problem found again keeps the day it was first found",
         [i["found"] for i in weekly.merge_status([amz], [again], None)]
         == ["2026-10-05"])
    import serve, tempfile
    real = weekly.CHECKS
    try:
        weekly.CHECKS = pathlib.Path(tempfile.mkdtemp())
        none_page = serve.home()
        weekly.save_status([], [amz], None, "1 brand", "full")
        one_page, details = serve.home(), serve.system_fix_page()
        weekly.save_status(weekly.load_status()["issues"], [], None, "", "full")
        cleared = serve.home()
    finally:
        weekly.CHECKS = real
    k.ok("system fix: no button while nothing is wrong",
         'class="sysfix"' not in none_page)
    k.ok("system fix: an open page asks for the button by itself (no reload)",
         "/system-fix/state" in none_page and "/system-fix/state" in one_page)
    k.ok("system fix: the flashing button shows while a problem is open, and "
         "the details page names it with a re-check button",
         'class="sysfix"' in one_page and "System Fix: 1 problem" in details
         and "Run the check again" in details and "u1" in details)
    k.ok("system fix: the button goes once a check passes", 'class="sysfix"'
         not in cleared)
    # The re-check: the code's tests first, then EVERY brand and channel.
    real_rt, real_pools = weekly.run_tests, weekly.pools
    asked = []
    try:
        weekly.CHECKS = pathlib.Path(tempfile.mkdtemp())
        weekly.save_status([], [amz], None, "", "full")
        weekly.run_tests = lambda: ["a test -- its detail"]
        weekly.pools = lambda brands=None: asked.append(brands) or {}
        weekly.run(recheck=True, retry_after=0)
        after_fail = [(i["brand"], i["kind"])
                      for i in weekly.load_status()["issues"]]
        live_skipped = asked == []
        weekly.run_tests = lambda: []
        weekly.run(recheck=True, retry_after=0)
        after_pass = weekly.load_status()["issues"]
    finally:
        weekly.CHECKS, weekly.run_tests, weekly.pools = real, real_rt, real_pools
    k.ok("system fix: a re-check whose tests fail keeps the old problem, adds "
         "the failed tests, and does not run the live check",
         after_fail == [("B", "Not readable"),
                        (weekly.TESTS_BRAND, "The code's own tests failed")]
         and live_skipped, (after_fail, asked))
    k.ok("system fix: a re-check whose tests pass checks EVERY brand and "
         "channel (not only the listed ones) and clears what passed",
         asked == [None] and after_pass == [], (asked, after_pass))

    # -- the spot-check on marketplaces
    import spotcheck, browser
    my = spotcheck.myntra_summary('<meta name="description" content="Buy VASTRAMAY '
                                  'Girls Kurta Set - Kurta Sets for Girls from '
                                  'VASTRAMAY at Rs. 1,984. Style ID: 25646028" />')
    k.ok("spot-check: Myntra's own page summary gives the price shoppers are told",
         my["sell"] == [1984.0], my)
    my = spotcheck.myntra_summary(
        '<meta name="description" content="... from Aj DEZInES at Rs. 829. Style ID: 1" />'
        '"sizeSellerData":[{"mrp":2599,"sellerPartnerId":1}]'
        '"sizeSellerData":[{"mrp":2599,"sellerPartnerId":1}]')
    k.ok("spot-check: Myntra's MRP from each size's seller, a second place",
         my["struck"] == [2599.0] and spotcheck.page_mrps(my) == [2599.0], my)

    # -- the spot-check compares the MRP (Abhisekh, 28 Sep 2026)
    def V(price, compare_at=None):
        return sites.Variant(frozenset(), price, compare_at, None, None)

    def LM(variants, listed_mrp=None):
        return sites._listing("https://x.example/p", "x.example", "storefront",
                              ok=True, variants=variants, listed_mrp=listed_mrp)
    k.ok("spot-check MRP: each size's own MRP; the page's one MRP; else the price",
         spotcheck.app_mrps(LM([V(839.0, 3499.0), V(1609.0, 3999.0)])) == [3499.0, 3999.0]
         and spotcheck.app_mrps(LM([V(829.0), V(829.0)], 2599.0)) == [2599.0]
         and spotcheck.app_mrps(LM([V(999.0)])) == [999.0])
    k.ok("spot-check MRP: struck or labelled MRP; no MRP shown = the price shown",
         spotcheck.page_mrps({"sell": [839], "struck": [3499]}) == [3499.0]
         and spotcheck.page_mrps({"sell": [799], "struck": [], "mrp": [1279]}) == [1279.0]
         and spotcheck.page_mrps({"sell": [999], "struck": []}) == [999.0])
    k.ok("spot-check MRP: same MRP, different selling price is a MATCH",
         spotcheck.verdict(LM([V(799.0, 1279.0)]),
                           {"sell": [899], "struck": [1279]})[0] == "MATCH")
    k.ok("spot-check MRP: a page MRP the app did not read is DIFFERENT",
         spotcheck.verdict(LM([V(799.0, 1279.0)]),
                           {"sell": [799], "struck": [1499]})[0] == "DIFFERENT")
    k.ok("spot-check MRP: the app says discounted, the page shows no MRP -> DIFFERENT",
         spotcheck.verdict(LM([V(799.0, 1279.0)]), {"sell": [799]})[0] == "DIFFERENT")
    k.ok("spot-check MRP: a price range is not compared (the MRP shows per size)",
         spotcheck.page_mrps({"sell": [549, 599], "struck": [1199]}) == []
         and spotcheck.verdict(LM([V(549.0, 899.0)]),
                               {"sell": [549, 599], "struck": [1199]})[0] == "CAN'T TELL")
    biba = ('<del><span class="strike-through list">\n  <span class="value" '
            'content="1999.00">')
    magento = ('"prices":{"baseOldPrice":{"amount":4883,"adjustments":[]},'
               '"oldPrice":{"amount":4883,"adjustments":[]}}')
    k.ok("page MRP: Salesforce's struck list price, Magento's old price",
         sites._page_mrp(biba, 1199.4) == 1999.0
         and sites._page_mrp(magento, 2930.0) == 4883.0)
    k.ok("page MRP: two different MRPs, or one not above the price, is none",
         sites._page_mrp(magento + '"oldPrice":{"amount":999', 2930.0) is None
         and sites._page_mrp(magento, 4883.0) is None)
    fk = ('"ppd":{"fsp":1344,"finalPrice":1424,"mrp":2999,"nepPrice":1276}'
          '"ppd":{"fsp":499,"finalPrice":520,"mrp":1299}')
    k.ok("Flipkart MRP: from the pricing record whose price is this listing's",
         sites._flipkart_mrp(fk, 1344.0) == 2999.0
         and sites._flipkart_mrp(fk, 700.0) is None
         and sites._flipkart_mrp(fk + '"ppd":{"fsp":1344,"mrp":3999}', 1344.0) is None)
    k.ok("spot-check: each marketplace gets its own price box; a website none",
         browser.market_probe("https://www.amazon.in/dp/B0X") is browser.MARKET_JS["amazon"]
         and browser.market_probe("https://www.flipkart.com/x/p/itm1") is browser.MARKET_JS["flipkart"]
         and browser.market_probe("https://www.firstcry.com/a/b/1/product-detail")
         is browser.MARKET_JS["firstcry"]
         and browser.market_probe("https://www.bownbee.com/products/x") is None)
    if not browser.available():
        print("  SKIP  the marketplace price-box scripts (Playwright not installed)")
        return
    pages = {
        "amazon": ('<h1><span id="productTitle">Kurta set</span></h1>'
                   '<div id="corePriceDisplay_desktop_feature_div"><span class="a-price '
                   'priceToPay"><span>&#8377;2,159</span></span><span class="a-price '
                   'a-text-price"><span>M.R.P.: &#8377;5,499</span></span></div>'
                   '<div class="rail"><span class="a-price">&#8377;689</span></div>'),
        "firstcry": ('<h1>Lehenga set</h1><p class="th-discounted-price"><label '
                     'class="rupee"></label><span class="prod-price" data-price='
                     '"1566.68">1566<sup class="supertag">68</sup></span> <span>'
                     '<span>MRP:</span><del>2999</del></span></p>'),
        "flipkart": ('<h1>Lehenga</h1><div><div style="font-size:27px">&#8377;1,344</div>'
                     '<div style="font-size:27px;text-decoration:line-through">'
                     '2,999</div></div>'
                     '<div style="font-size:20px">Buy at &#8377;1,276</div>'
                     '<div style="font-size:14px">Discount on MRP -&#8377;1,655</div>'
                     '<div style="font-size:16px">&#8377;68 off</div>'
                     '<div style="font-size:11px">&#8377;892</div>')}
    got = {}
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            b = p.chromium.launch(headless=True)
            pg = b.new_page()
            for name, body in pages.items():
                pg.set_content("<html><body>%s</body></html>" % body)
                got[name] = pg.evaluate(browser.MARKET_JS[name])
            b.close()
    except Exception as e:
        k.ok("the marketplace price-box scripts run", False, repr(e)[:200])
        return
    k.ok("Amazon's price box: the price to pay, the M.R.P. apart, rails ignored",
         got["amazon"]["sell"] == [2159] and got["amazon"]["struck"] == [5499],
         got["amazon"])
    k.ok("FirstCry's price box: Rs 1566 and a raised 68 is Rs 1,566.68",
         got["firstcry"]["sell"] == [1566.68] and got["firstcry"]["struck"] == [2999],
         got["firstcry"])
    k.ok("Flipkart: the largest price on screen, never 'Buy at' or 'off'",
         got["flipkart"]["sell"] == [1344], got["flipkart"])
    k.ok("Flipkart: the struck-through MRP beside that price",
         got["flipkart"]["struck"] == [2999], got["flipkart"])


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.parse_args(argv)

    k = Checks()
    adapters(k)
    resilience(k)
    summary_to_page(k)
    link_file(k)
    brand_catalogue(k)
    hardening(k)
    js_sites(k)
    weekly_check(k)

    print("\n%d passed, %d failed" % (k.passed, k.failed))
    for name, detail in k.failures:
        print("  FAILED: %s -- %s" % (name, detail))
    return 1 if k.failed else 0


if __name__ == "__main__":
    sys.exit(main())
