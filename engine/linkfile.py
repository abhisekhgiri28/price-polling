"""Read the price-polling input file: a barcode, its design, and its links.

This is the ONLY input the app takes (the ZOCS intake CSV it used to take is
retired from the page). One row per barcode:

    zoddle_barcode | design_code | size | zoddle_price | image | Myntra | FirstCry | Amazon | Flipkart | Website

  * Required columns: zoddle barcode, design code, size, zoddle price, image.
  * Plus at least one marketplace column. The HEADING names the marketplace
    ("Myntra", "FirstCry", "Amazon", "Flipkart", "Website" / "Own site" /
    "<Brand> Website", "Meesho", ...) and each cell is that marketplace's
    product-page link for this barcode. A blank cell means no link.
  * Optional: mrp, design name, brand, colour, brand size.

.csv and .xlsx are both accepted. In an .xlsx, a cell's hyperlink wins over
its display text, because a sheet often shows "Myntra" and links the page.

Every link is checked against its heading. A Myntra URL sitting under
"Amazon" is dropped with a finding rather than priced as Amazon -- a price
must be traceable to the link AND labelled with the channel it came from.

Returns the same shape as intake.parse_bytes -- (records, issues, meta,
encoding) -- so compare.run takes either. Each record carries `links`, the
(channel, url) pairs from the named columns, which compare.channels_for uses
in preference to guessing a channel from a URL's host.
"""
import csv, io, re, collections, urllib.parse

# Same shape as intake.Issue, so the page renders both without caring.
Issue = collections.namedtuple("Issue", "row_no barcode severity code detail")

# Heading -> field. Headings are compared lower-case with punctuation folded
# to single spaces, so "Zoddle Barcode", "zoddle_barcode" and "ZODDLE-BARCODE"
# are one heading.
FIELDS = {
    "zoddle_barcode": ("zoddle barcode", "barcode", "zoddle bar code"),
    "design_code": ("design code", "design id", "zoddle design code"),
    "size": ("size", "zoddle size"),
    "zoddle_price": ("zoddle price", "zoddle selling price"),
    "image": ("image", "image link", "image url", "image url 1",
              "zoddle image", "photo", "photo link", "picture"),
    "mrp": ("mrp", "zoddle mrp"),
    "name": ("design name", "product name", "name", "design"),
    "brand": ("brand", "brand name"),
    "color": ("color", "colour"),
    "brand_size": ("brand size",),
}
REQUIRED = ("zoddle_barcode", "design_code", "size", "zoddle_price", "image")
LABEL = {"zoddle_barcode": "Zoddle barcode", "design_code": "Design code",
         "size": "Size", "zoddle_price": "Zoddle price", "image": "Image link"}

# Marketplace column headings. Checked in order; the first keyword found in a
# heading names its channel. `hosts` is what that channel's links must be on.
MARKETPLACES = (
    ("myntra", ("myntra",), ("myntra.com",)),
    ("firstcry", ("firstcry", "first cry"), ("firstcry.com",)),
    ("amazon", ("amazon",), ("amazon.in", "amazon.com", "amzn.in", "amzn.to")),
    ("flipkart", ("flipkart",), ("flipkart.com",)),
    ("meesho", ("meesho",), ("meesho.com",)),
    ("tatacliq", ("tatacliq", "tata cliq"), ("tatacliq.com",)),
    ("ajio", ("ajio",), ("ajio.com",)),
    ("nykaa", ("nykaa",), ("nykaa.com", "nykaafashion.com")),
    ("hopscotch", ("hopscotch",), ("hopscotch.in",)),
)
NOT_LINK_WORDS = ("price", "mrp", "rs", "inr", "status", "size", "sizes",
                  "colour", "color", "sku", "asin", "fsn", "id", "code",
                  "stock", "qty", "quantity", "discount", "rating", "date",
                  "remark", "remarks", "note", "notes", "comment", "comments")
OWN_SITE_WORDS = ("website", "own site", "brand site", "web site", "site",
                  "store", "shop", "d2c")
MARKETPLACE_HOSTS = tuple(h for _, _, hs in MARKETPLACES for h in hs)


def _fold(s):
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()


def _host(url):
    return urllib.parse.urlsplit(url).netloc.lower()


def _on(host, hosts):
    return any(host == h or host.endswith("." + h) for h in hosts)


def channel_for_heading(heading):
    """'Myntra link' -> 'myntra'; 'Vastramay Website' -> 'own site'; else None."""
    h = _fold(heading)
    if not h:
        return None
    # "Myntra price", "Myntra ₹", "Amazon MRP" are figures, not links; and
    # "Myntra size", "Amazon ASIN", "FirstCry style id" name a marketplace
    # but hold something else -- read as links, their cells were skipped
    # and the download overwrote them with "-".
    if "₹" in str(heading) or any(w in h.split() for w in NOT_LINK_WORDS):
        return None
    for name, words, _hosts in MARKETPLACES:
        if any(w in h for w in words):
            return name
    words = h.split()
    if any((" " in w and w in h) or w in words for w in OWN_SITE_WORDS):
        return "own site"
    return None


# A web address with no scheme: "www." in front, or a path after the host.
# "Not.Available" or "Size.XL" are words, not links.
_BARE_HOST = re.compile(r"^(?:www\.[a-z0-9-]+(?:\.[a-z0-9-]+)*\.[a-z]{2,}(?:[/?#]|$)"
                        r"|[a-z0-9-]+(?:\.[a-z0-9-]+)*\.[a-z]{2,}[/?#])", re.I)


def _with_scheme(url):
    """'www.myntra.com/x/1/buy' -> 'https://www.myntra.com/x/1/buy'. A link
    copied from the address bar often loses its https://, and was skipped."""
    if url.lower().startswith(("http://", "https://")):
        return url
    if url.startswith("//"):
        return "https:" + url
    return "https://" + url if _BARE_HOST.match(url) else url


def is_link(cell):
    """Does this cell hold a web link (with or without https://)?"""
    return _with_scheme((cell or "").strip()).lower().startswith(
        ("http://", "https://"))


def _money(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v) if v > 0 else None
    # "1,299", "Rs. 1299", "₹1,299.00", "INR 1299" and "1299/-" are all
    # written in real files; each used to reject the row as having no price.
    t = str(v).strip().replace(",", "").replace("₹", "")
    t = re.sub(r"^(?:rs\.?|inr)\s*", "", t, flags=re.I)
    t = re.sub(r"\s*/-$", "", t).strip()
    try:
        f = float(t)
    except ValueError:
        return None
    return f if f > 0 else None


def _text(v):
    """A cell as clean text. Excel hands barcodes back as 4470134.0."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def size_code(size):
    """'2-3Y' -> '23', '6-12M' -> '62'. None when the size has one bound."""
    d = re.findall(r"\d+", size or "")
    return (d[0][-1] + d[1][-1]) if len(d) >= 2 else None


# ------------------------------------------------------------------ reading --
def _rows_from_csv(raw):
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
        except UnicodeDecodeError:
            continue
        return list(csv.reader(io.StringIO(text, newline=""),
                               delimiter=_delimiter(text))), enc
    return None, None


def _delimiter(text):
    """Excel saves "CSV" with a semicolon in some regional settings, and a
    copy-paste gives tabs; a header line with neither comma is read by what
    it does contain rather than rejected as one column."""
    head = [ln for ln in text.splitlines() if ln.strip()][:10]
    counts = {d: sum(ln.count(d) for ln in head) for d in (",", ";", "\t")}
    best = max(counts, key=lambda d: (counts[d], d == ","))
    return best if counts[best] else ","


def _rows_from_xlsx(raw):
    """First sheet, as rows of text. A cell's hyperlink wins over its text."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.worksheets[0]
    # The same sheet with its formulas, for =HYPERLINK("url", "Myntra"):
    # its value is only the label (or nothing, if never opened in Excel),
    # so the link used to be dropped -- silently, when there was no value.
    formulas = openpyxl.load_workbook(io.BytesIO(raw)).worksheets[0]
    out = []
    for row in ws.iter_rows():
        vals = []
        for c in row:
            link = c.hyperlink.target if c.hyperlink is not None else None
            if not link:
                link = _hyperlink_formula(formulas.cell(row=c.row,
                                                        column=c.column).value)
            vals.append(link if link else _text(c.value))
        out.append(vals)
    return out, "xlsx"


_HYPERLINK = re.compile(r'^=\s*HYPERLINK\(\s*"([^"]+)"', re.I)


def _hyperlink_formula(value):
    """'=HYPERLINK("https://x", "Myntra")' -> 'https://x', else None."""
    m = _HYPERLINK.match(value) if isinstance(value, str) else None
    return m.group(1) if m else None


def is_xlsx(raw, filename=""):
    return filename.lower().endswith((".xlsx", ".xlsm")) or raw[:2] == b"PK"


def parse_bytes(raw, filename=""):
    """(records, issues, meta, encoding). Nothing is written, nothing raises."""
    issues = []
    try:
        if is_xlsx(raw, filename):
            table, enc = _rows_from_xlsx(raw)
        else:
            table, enc = _rows_from_csv(raw)
    except Exception as ex:
        return [], [Issue(0, None, "reject", "unreadable",
                          "could not read the file (%s)" % type(ex).__name__)], \
            {"filename": filename}, None
    if not table:
        return [], [Issue(0, None, "reject", "empty_file",
                          "the file is empty or not readable text")], \
            {"filename": filename}, enc
    records, issues, meta = parse_table(table, filename)
    return records, issues, meta, enc


def layout(table):
    """(head_at, header, col, links) -- where the header row is, its cells,
    field -> column index, and (column index, channel, heading) per link
    column. Shared by the reader and by priced_csv, so the file that comes
    back is laid out by exactly the rules the file that went in was read by.
    """
    # The header is the first row that names a barcode column, so a title
    # line above the table does not break the file.
    head_at = next((i for i, r in enumerate(table[:10])
                    if any(_fold(c) in FIELDS["zoddle_barcode"] for c in r)), 0)
    header = [_text(c) for c in table[head_at]]
    folded = [_fold(h) for h in header]

    col = {}
    for field, names in FIELDS.items():
        for i, h in enumerate(folded):
            if h in names and i not in col.values():
                col[field] = i
                break
    links = []                                   # (column index, channel, heading)
    for i, h in enumerate(header):
        if i in col.values():
            continue
        ch = channel_for_heading(h)
        if ch:
            links.append((i, ch, h))
    return head_at, header, col, links


def read_table(raw, filename=""):
    """The file's rows of text, header included, or None if unreadable."""
    try:
        table, _ = (_rows_from_xlsx(raw) if is_xlsx(raw, filename)
                    else _rows_from_csv(raw))
    except Exception:
        return None
    return table


def _price_text(v):
    if v in (None, ""):
        return ""
    v = float(v)
    return "%d" % v if abs(v - round(v)) < 0.005 else "%.2f" % v


NA = "-"          # a price that is not available, in the CSV download


class Kept(str):
    """A link column's cell that held no link, handed back as written. It
    is never stored as a number in the Excel file: a "1299" typed under
    Myntra was never fetched, so it must not look like a live price."""
BLANK_LINKS = ("-", "—", "–", "NA", "N/A", "na", "n/a")


def priced_csv(table, prices, lowest=None, drop_headings=("last_updated",)):
    """The priced file as CSV text -- see priced_rows for the layout."""
    head, rows, _ = priced_rows(table, prices, lowest, drop_headings)
    buf = io.StringIO(newline="")
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(head)
    w.writerows(rows)
    return buf.getvalue()


def priced_xlsx(table, prices, lowest=None, drop_headings=("last_updated",),
                links=None):
    """The priced file as an Excel workbook (bytes): the same rows as the
    CSV, with prices stored as numbers and every "-" centred in its cell
    (Abhisekh, 24 Sep 2026 -- a CSV cannot carry alignment).

    Every price -- and "out of stock" / "not fetched" -- is a link to the
    page it was read from (Abhisekh, 30 Sep 2026), as on the app's page:
    `links` maps (barcode, channel) -> that page, and `lowest` may carry
    the Lowest price's page as a third item. A CSV cannot carry links."""
    import openpyxl
    from openpyxl.styles import Alignment, Font
    head, rows, numeric, hrefs = _priced(table, prices, lowest,
                                         drop_headings, links)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Live prices"
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    centre = Alignment(horizontal="center")
    def num(v):
        m = _money(v)                   # "1,299" / "Rs 1299" -> 1299.0
        return v if m is None else (int(m) if m == int(m) else m)

    for n, r in enumerate(rows):
        ws.append([num(v) if j in numeric and v != NA
                   and not isinstance(v, Kept) else str(v)
                   if isinstance(v, Kept) else v
                   for j, v in enumerate(r)])
        for j, c in enumerate(ws[ws.max_row]):
            if c.value == NA:
                c.alignment = centre
            elif (n, j) in hrefs:
                c.hyperlink = hrefs[(n, j)]
                c.style = "Hyperlink"
    for j, h in enumerate(head, start=1):
        width = max([len(str(h))] + [len(str(r[j - 1])) for r in rows])
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = \
            min(40, width + 2)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def priced_rows(table, prices, lowest=None, drop_headings=("last_updated",)):
    """(header, rows, numeric column positions): the input file back, with
    every marketplace link replaced by the live price read from it, and the
    cheapest price and its source added.

    Laid out as Abhisekh asked (24 Sep 2026): design code, Zoddle barcode,
    size, MRP, Zoddle price, then the file's other columns in their order
    (the marketplace columns among them), then `Lowest price` and `Lowest
    source`. The image, design name and colour columns -- and any heading in
    `drop_headings` -- are left out. Every price cell with no price, whether
    the link gave none or there was no link, reads "-".

    `prices` maps (barcode, channel) -> price; `lowest` maps barcode ->
    (price, source name).
    """
    return _priced(table, prices, lowest, drop_headings)[:3]


def _priced(table, prices, lowest=None, drop_headings=("last_updated",),
            page_links=None):
    """priced_rows, plus {(output row, column): page} for every cell that
    shows a price or a word, from `page_links` ((barcode, channel) -> page)
    and the third item of a `lowest` entry, when given."""
    head_at, header, col, links = layout(table)
    page_links = page_links or {}
    hrefs = {}
    lowest = lowest or {}
    drop = {col[f] for f in ("image", "name", "color") if f in col}
    drop |= {i for i, h in enumerate(header)
             if _fold(h) in {_fold(d) for d in drop_headings}}
    link_ch = {i: ch for i, ch, _ in links}
    # A column no row fills (a saved brand's brand_size, say) is noise.
    body = [[_text(c) for c in r] for r in table[head_at + 1:]]
    drop |= {i for i in range(len(header)) if i not in link_ch
             and i not in col.values()
             and not any(i < len(r) and r[i].strip() for r in body)}
    drop |= {col["brand_size"]} if "brand_size" in col and not any(
        col["brand_size"] < len(r) and r[col["brand_size"]].strip()
        for r in body) else set()
    lead = [col[f] for f in ("design_code", "zoddle_barcode", "size", "mrp",
                             "zoddle_price") if f in col]
    keep = lead + [i for i in range(len(header))
                   if i not in drop and i not in lead]
    money = {col[f] for f in ("mrp", "zoddle_price") if f in col}
    bc = col.get("zoddle_barcode")

    head = [header[i] for i in keep] + ["Lowest price", "Lowest source"]
    # Positions in the OUTPUT row that hold a price, for the Excel file.
    numeric = {j for j, i in enumerate(keep) if i in link_ch or i in money}
    numeric.add(len(keep))
    out = []
    # A row the reader refused (a repeated barcode, no Zoddle price ...)
    # was never fetched, so it gets no prices -- whichever of two rows with
    # one barcode that is. Row numbers as parse_table counts them.
    refused = {i.row_no for i in parse_table(table)[1]
               if i.severity == "reject"}
    for n, row in enumerate(table[head_at + 1:], start=head_at + 2):
        cells = [_text(c) for c in row]
        if not any(cells):
            continue
        cells += [""] * (len(header) - len(cells))
        barcode = cells[bc].strip() if bc is not None else ""
        repeat = n in refused

        def cell(i):
            if i in link_ch:
                if cells[i].strip() and not is_link(cells[i]) and \
                        cells[i].strip() not in BLANK_LINKS:
                    return Kept(cells[i])      # not a link: leave it be
                # A price only where the file had a link.
                got = (prices.get((barcode, link_ch[i]))
                       if cells[i].strip() and not repeat else "")
                if isinstance(got, str) and got and _money(got) is None:
                    return Kept(got)           # "out of stock", "not fetched"
                return _price_text(got) or NA
            if i in money:
                return cells[i] or NA
            return cells[i]
        low = ("", "") if repeat else tuple(lowest.get(barcode, ("", "")))
        low_price, low_source = low[:2]
        line = ([cell(i) for i in keep]
                + [_price_text(low_price) or NA, low_source or NA])
        for j, i in enumerate(keep):
            page = (page_links.get((barcode, link_ch[i]))
                    if i in link_ch else None)
            if page and line[j] != NA and cells[i].strip() and not repeat \
                    and is_link(cells[i]):
                hrefs[(len(out), j)] = page
        if len(low) > 2 and low[2] and line[len(keep)] != NA:
            hrefs[(len(out), len(keep))] = low[2]
        out.append(line)
    return head, out, numeric, hrefs


def parse_table(table, filename=""):
    """Rows of text (header first) -> (records, issues, meta)."""
    issues = []
    head_at, header, col, links = layout(table)
    ignored = [h for i, h in enumerate(header)
               if h and i not in col.values() and i not in [c for c, _, _ in links]]

    meta = {"filename": filename, "header": header,
            "link_columns": [(h, ch) for _, ch, h in links],
            "ignored_columns": ignored, "rows_read": 0}

    missing = [LABEL[f] for f in REQUIRED if f not in col]
    if missing or not links:
        why = []
        if missing:
            why.append("missing required column%s: %s"
                       % ("" if len(missing) == 1 else "s", ", ".join(missing)))
        if not links:
            why.append("no marketplace column -- name each link column after "
                       "its marketplace (Myntra, FirstCry, Amazon, Flipkart, "
                       "Website ...)")
        issues.append(Issue(1, None, "reject", "bad_header",
                            "; ".join(why) + ". Columns found: "
                            + (", ".join(h for h in header if h) or "none")))
        return [], issues, meta

    seen = {}
    records = []
    for n, row in enumerate(table[head_at + 1:], start=head_at + 2):
        cells = [_text(c) for c in row]
        if not any(cells):
            continue
        meta["rows_read"] += 1

        def get(field):
            i = col.get(field)
            return cells[i].strip() if i is not None and i < len(cells) else ""

        barcode = get("zoddle_barcode")
        bad = []

        def add(sev, code, detail):
            issues.append(Issue(n, barcode or None, sev, code, detail))
            if sev == "reject":
                bad.append(code)

        if not barcode:
            add("reject", "no_barcode", "Zoddle barcode is blank")
        elif not barcode.isdigit():
            add("reject", "barcode_not_numeric", "Zoddle barcode %r" % barcode)
        elif len(barcode) < 5:
            add("reject", "barcode_too_short",
                "%r is shorter than design + colourway + size" % barcode)
        elif barcode in seen:
            add("reject", "duplicate_barcode",
                "already on row %d of this file" % seen[barcode])

        design = get("design_code")
        if not design:
            add("reject", "no_design_code", "Design code is blank")
        elif barcode.isdigit() and len(barcode) >= 5 and design != barcode[:-4]:
            add("reject", "design_code_conflict",
                "design code %s, but barcode %s belongs to design %s"
                % (design, barcode, barcode[:-4]))

        size = get("size")
        if not size:
            add("reject", "no_size", "Size is blank -- a price is per size")
        elif barcode[-2:].isdigit() and size_code(size) and \
                size_code(size) != barcode[-2:]:
            add("warn", "size_barcode_mismatch",
                "barcode ends %s but size %r implies %s"
                % (barcode[-2:], size, size_code(size)))

        price = _money(get("zoddle_price"))
        if price is None:
            add("reject", "no_zoddle_price",
                "Zoddle price %r is not a positive number" % get("zoddle_price"))

        image = get("image")
        if not image:
            add("warn", "no_image", "no image link -- the thumbnail will show "
                                    "only the marketplace's photograph")
        elif not image.lower().startswith(("http://", "https://")):
            add("warn", "image_not_a_link", "image %r is not a web link" % image)
            image = ""

        row_links = []
        for i, ch, heading in links:
            url = cells[i].strip() if i < len(cells) else ""
            if not url or url in BLANK_LINKS:
                continue
            url = _with_scheme(url)
            if not url.lower().startswith(("http://", "https://")):
                add("warn", "link_not_a_url",
                    "%s: %r is not a web link, so it was skipped" % (heading, url))
                continue
            host = _host(url)
            want = next((hs for name, _, hs in MARKETPLACES if name == ch), None)
            if want and not _on(host, want):
                add("warn", "link_wrong_marketplace",
                    "%s column holds a %s link, so it was skipped" % (heading, host))
                continue
            if ch == "own site" and _on(host, MARKETPLACE_HOSTS):
                add("warn", "link_wrong_marketplace",
                    "%s column holds a marketplace link (%s), so it was "
                    "skipped" % (heading, host))
                continue
            row_links.append((ch, url))
        if not row_links and not bad:
            add("note", "no_links", "no marketplace link on this row -- it "
                                    "will be listed with no prices")

        if bad:
            continue
        seen[barcode] = n
        rec = {
            "zoddle_barcode": barcode, "design_code": design,
            "colorway_code": barcode[-4:-2], "size_code": barcode[-2:],
            "size": size, "zoddle_price": price, "mrp": _money(get("mrp")),
            "name": get("name") or None, "brand": get("brand") or None,
            "color": get("color") or None, "brand_size": get("brand_size") or None,
            "images": ([{"position": 1, "url": image, "host": _host(image),
                         "is_page": False}] if image else []),
            "links": row_links,
        }
        records.append(rec)

    meta["rows_loaded"] = len(records)
    brands = collections.Counter(r["brand"] for r in records if r["brand"])
    meta["brands"] = [b for b, _ in brands.most_common()]
    return records, issues, meta


# ------------------------------------------------------------------ template --
TEMPLATE_HEADER = ["zoddle_barcode", "design_code", "size", "zoddle_price",
                   "mrp", "design_name", "image", "Website", "Myntra",
                   "FirstCry", "Amazon", "Flipkart"]


def template_csv():
    """An empty input file with the headings filled in, plus one example row."""
    buf = io.StringIO(newline="")
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(TEMPLATE_HEADER)
    w.writerow(["4470134", "447", "3-4Y", "1114", "3599",
                "Girls Lehenga Choli", "https://example.com/photo.jpg",
                "https://brand.com/products/x", "https://www.myntra.com/x/1/buy",
                "", "https://www.amazon.in/dp/B0XXXXXXX", ""])
    return buf.getvalue()
