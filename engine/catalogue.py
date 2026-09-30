"""The brand catalogue: every uploaded product, kept so it can be re-polled.

One Excel workbook, `data/brand_catalogue.xlsx`, with one sheet per brand.
Each sheet holds one row per Zoddle barcode -- the same columns as the input
file (see linkfile.py) plus `last_updated` -- sorted by design code, then
barcode. Excel rather than CSV because a CSV cannot hold a sheet per brand.

Merging an upload into a brand's sheet (Abhisekh's rules, 23 Sep 2026):

  * A design code already in the sheet: its marketplace links are updated on
    EVERY barcode of that design. For each marketplace, a link in the new
    file replaces the old one; a marketplace the new file leaves blank keeps
    the link it already had.
  * A new barcode of an existing design is added, and inherits the design's
    links (then the new file's, per the rule above).
  * A new design code is added with all its barcodes.
  * For a barcode already present, the other fields (size, Zoddle price,
    MRP, name, colour, image) take the new file's value where it gives one,
    and keep the old value where it is blank.

The workbook is rewritten atomically (a temporary file, then a rename), so a
crash mid-save cannot leave it half-written. If Excel has it open, Windows
refuses the rename; that is reported as such and nothing is lost.
"""
import os, re, io, datetime, threading, collections, pathlib, tempfile

import linkfile

HERE = pathlib.Path(__file__).resolve().parent
PATH = HERE.parent / "data" / "brand_catalogue.xlsx"
_lock = threading.Lock()

# brand_size is the brand's own name for the size ("1-2Y" where Zoddle says
# "12-18M"); a page is matched on it first. It was read from uploads but not
# saved, so a saved brand lost it (found 26 Sep 2026).
BASE = ["zoddle_barcode", "design_code", "size", "brand_size", "zoddle_price",
        "mrp", "design_name", "colour", "image"]
# The heading each channel is written under, in the order shown.
HEADING = collections.OrderedDict([
    ("own site", "Website"), ("myntra", "Myntra"), ("firstcry", "FirstCry"),
    ("amazon", "Amazon"), ("flipkart", "Flipkart"), ("meesho", "Meesho"),
    ("tatacliq", "TataCliq"), ("ajio", "Ajio"), ("nykaa", "Nykaa"),
    ("hopscotch", "Hopscotch")])
TAIL = ["last_updated"]


class CatalogueLocked(Exception):
    """The workbook could not be written -- almost always open in Excel."""


def sheet_name(brand):
    """Excel sheet names: at most 31 characters, none of []:*?/\\. A longer
    name is cut and given a short mark of the whole name, so two brands
    alike in their first 31 letters do not share one sheet."""
    name = re.sub(r"[\[\]:*?/\\]", " ", str(brand or "").strip())
    name = re.sub(r"\s+", " ", name).strip()
    if len(name) > 31:
        import zlib
        mark = "%04x" % (zlib.crc32(name.casefold().encode("utf-8")) & 0xffff)
        name = name[:25].rstrip() + " ~" + mark
    return name or "Unbranded"


def _sheet_for(wb, brand):
    """The brand's existing sheet title, matched ignoring case, or None.

    Excel treats sheet names case-insensitively, and openpyxl renames a
    clashing new sheet to "bownbee1" -- so an upload typed "bownbee" used to
    start a second, orphan sheet beside "BownBee" instead of merging.
    """
    want = sheet_name(brand).casefold()
    return next((t for t in wb.sheetnames if t.casefold() == want), None)


def _design_key(code):
    code = str(code or "")
    return (0, int(code), "") if code.isdigit() else (1, 0, code)


def _row_key(row):
    return (_design_key(row.get("design_code")), str(row.get("zoddle_barcode")))


def _money_text(v):
    if v in (None, ""):
        return ""
    f = float(v)
    return int(f) if f.is_integer() else round(f, 2)


# ------------------------------------------------------------------- reading --
def _load(path=None):
    import openpyxl
    path = pathlib.Path(path or PATH)
    if not path.exists():
        return openpyxl.Workbook(), False
    return openpyxl.load_workbook(path), True


def _sheet_rows(ws):
    """A sheet as a list of dicts keyed by heading. Hyperlinks win over text."""
    rows = list(ws.iter_rows())
    if not rows:
        return [], []
    header = [linkfile._text(c.value) for c in rows[0]]
    out = []
    for r in rows[1:]:
        d = {}
        for h, c in zip(header, r):
            if not h:
                continue
            link = c.hyperlink.target if c.hyperlink is not None else None
            d[h] = link or linkfile._text(c.value)
        if d.get("zoddle_barcode"):
            out.append(d)
    return header, out


def brands(path=None):
    """[(brand, designs, barcodes)] for the dropdown, alphabetical."""
    with _lock:
        wb, exists = _load(path)
        if not exists:
            return []
        out = []
        for ws in wb.worksheets:
            _, rows = _sheet_rows(ws)
            if rows:
                out.append((ws.title, len({r.get("design_code") for r in rows}),
                            len(rows)))
        return sorted(out, key=lambda t: t[0].lower())


def designs(brand, path=None):
    """The design codes stored for one brand, in sheet order."""
    with _lock:
        wb, exists = _load(path)
        name = _sheet_for(wb, brand) if exists else None
        if not name:
            return []
        _, rows = _sheet_rows(wb[name])
        return list(dict.fromkeys(r.get("design_code") for r in rows))


def saved_name(brand, path=None):
    """The brand as its sheet is actually titled ("bownbee" -> "BownBee"),
    or the sheet name it would get if it is not saved yet."""
    with _lock:
        wb, exists = _load(path)
        return (_sheet_for(wb, brand) if exists else None) or sheet_name(brand)


def table_for(brand, design_codes=None, path=None):
    """A stored brand as rows of text, header first -- the same shape as an
    uploaded file -- or None when the brand is not saved. `design_codes`
    narrows it to those designs."""
    with _lock:
        wb, exists = _load(path)
        name = _sheet_for(wb, brand) if exists else None
        if not name:
            return None
        header, rows = _sheet_rows(wb[name])
    if design_codes:
        want = {str(d).strip() for d in design_codes if str(d).strip()}
        rows = [r for r in rows if r.get("design_code") in want]
    return [header] + [[r.get(h, "") for h in header] for r in rows]


def records_for(brand, design_codes=None, path=None):
    """(records, issues, meta) for a stored brand -- ready for compare.run.

    `design_codes` narrows it to those designs; None means the whole brand.
    The sheet is read back through linkfile, so a stored row is checked by
    exactly the rules an uploaded one is.
    """
    name = saved_name(brand, path)
    table = table_for(brand, design_codes, path)
    if table is None:
        return [], [linkfile.Issue(0, None, "reject", "no_such_brand",
                                   "no saved products for %r" % brand)], {}
    records, issues, meta = linkfile.parse_table(table, name)
    for rec in records:
        rec["brand"] = rec.get("brand") or name
    return records, issues, meta


# ------------------------------------------------------------------- merging --
def merge(brand, records, path=None, today=None):
    """Fold uploaded records into the brand's sheet. Returns what changed.

    `records` are linkfile records (each with `links` as (channel, url)).
    """
    today = today or datetime.date.today().isoformat()
    stats = collections.Counter()
    with _lock:
        wb, exists = _load(path)
        # An existing sheet keeps its own spelling: "bownbee" merges into
        # "BownBee" rather than starting a second sheet.
        name = (_sheet_for(wb, brand) if exists else None) or sheet_name(brand)
        if name in wb.sheetnames:
            header, rows = _sheet_rows(wb[name])
        else:
            header, rows = [], []
        by_barcode = {r["zoddle_barcode"]: r for r in rows}
        by_design = collections.defaultdict(list)
        for r in rows:
            by_design[r.get("design_code")].append(r)

        # The new file's link per design and marketplace: the first
        # non-blank one among that design's rows.
        new_links = collections.defaultdict(collections.OrderedDict)
        grouped = collections.OrderedDict()
        for rec in records:
            grouped.setdefault(rec["design_code"], []).append(rec)
            for ch, url in rec.get("links") or []:
                h = HEADING.get(ch, ch.title())
                new_links[rec["design_code"]].setdefault(h, url)

        for design, recs in grouped.items():
            old_rows = by_design.get(design, [])
            if old_rows:
                stats["designs_updated"] += 1
            else:
                stats["designs_added"] += 1
            # The design's links as they stand, then the new file's on top.
            inherited = {}
            for r in old_rows:
                for h in HEADING.values():
                    if r.get(h) and h not in inherited:
                        inherited[h] = r[h]
                for h, v in r.items():
                    if h not in BASE + TAIL and v and h not in inherited \
                            and v.startswith("http"):
                        inherited[h] = v
            links = dict(inherited)
            for h, url in new_links[design].items():
                if links.get(h) and links[h] != url:
                    stats["links_replaced"] += 1
                elif not links.get(h):
                    stats["links_added"] += 1
                links[h] = url

            for rec in recs:
                row = by_barcode.get(rec["zoddle_barcode"])
                if row is None:
                    row = {"zoddle_barcode": rec["zoddle_barcode"],
                           "design_code": design}
                    by_barcode[rec["zoddle_barcode"]] = row
                    by_design[design].append(row)
                    stats["barcodes_added"] += 1
                else:
                    stats["barcodes_updated"] += 1
                image = next((i["url"] for i in rec.get("images") or []
                              if not i.get("is_page")), "")
                for field, value in (
                        ("size", rec.get("size")),
                        ("brand_size", rec.get("brand_size")),
                        ("zoddle_price", _money_text(rec.get("zoddle_price"))),
                        ("mrp", _money_text(rec.get("mrp"))),
                        ("design_name", rec.get("name")),
                        ("colour", rec.get("color")),
                        ("image", image)):
                    if value not in (None, ""):
                        row[field] = value
                row["last_updated"] = today

            # Every barcode of the design -- old and new -- carries the
            # design's links. A marketplace the file left blank kept its old
            # link above, so nothing is lost here.
            for row in by_design[design]:
                for h, url in links.items():
                    row[h] = url
                row.setdefault("last_updated", today)

        # Columns: the base fields, the marketplaces in their usual order,
        # anything else already on the sheet, then last_updated.
        present = {h for r in by_barcode.values() for h, v in r.items() if v}
        chans = [h for h in HEADING.values() if h in present]
        extra = [h for h in header if h and h not in BASE + TAIL + chans]
        extra += sorted(h for h in present
                        if h not in BASE + TAIL + chans + extra)
        columns = BASE + chans + extra + TAIL
        out_rows = sorted(by_barcode.values(), key=_row_key)

        if not exists and wb.sheetnames == ["Sheet"]:
            # A new workbook's own empty "Sheet": gone first, or a brand
            # called "sheet" is saved as "sheet1".
            del wb["Sheet"]
        if name in wb.sheetnames:
            idx = wb.sheetnames.index(name)
            del wb[name]
            ws = wb.create_sheet(name, idx)
        else:
            ws = wb.create_sheet(name)
        _write_sheet(ws, columns, out_rows)
        _save(wb, path)
        stats["barcodes_total"] = len(out_rows)
        stats["designs_total"] = len({r.get("design_code") for r in out_rows})
    return dict(stats)


def _write_sheet(ws, columns, rows):
    from openpyxl.styles import Font
    ws.append(columns)
    for c in ws[1]:
        c.font = Font(bold=True)
    link_cols = set(HEADING.values()) | {"image"}
    for r in rows:
        values = []
        for h in columns:
            v = r.get(h, "")
            if h in ("zoddle_price", "mrp") and v not in ("", None):
                try:
                    v = _money_text(v)
                except ValueError:
                    pass
            values.append(v)
        ws.append(values)
        for i, h in enumerate(columns, start=1):
            cell = ws.cell(row=ws.max_row, column=i)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"         # text, never a formula
            if (h in link_cols or h not in BASE + TAIL) and \
                    isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value
    ws.freeze_panes = "C2"
    widths = {"zoddle_barcode": 15, "design_code": 12, "size": 9,
              "brand_size": 11,
              "zoddle_price": 12, "mrp": 9, "design_name": 36, "colour": 12,
              "image": 30, "last_updated": 13}
    from openpyxl.utils import get_column_letter
    for i, h in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 34)


def _save(wb, path=None):
    path = pathlib.Path(path or PATH)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError:
        raise CatalogueLocked(
            "%s could not be saved -- it is probably open in Excel. Close it "
            "and upload again; nothing was changed." % path.name)
    finally:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass
